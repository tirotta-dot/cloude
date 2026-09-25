#!/usr/bin/env python3
"""Pagamenti aperti: il foglio settimanale per l'amministrazione e le sue risposte (R32, 25/09/2026).

Richiesta di Danilo: ogni settimana un foglio a Rossella Strati, Paola Barbieri ed Enrica Lorenzini con
i pagamenti che risultano aperti da contratto o che vanno fatturati; per ogni riga rispondono da un menu
a tendina (Pagato / Da sollecitare / In arrivo il ...), le risposte aggiornano il Second Brain e Danilo
riceve un report.

  # il foglio (xlsx con la tendina, da caricare su Drive convertito in Foglio Google) e l'elenco in JSON
  python3 pagamenti_aperti.py foglio --state state.json --out Pagamenti.xlsx --json voci.json [--oggi AAAA-MM-GG]

  # le risposte (raccolte dal foglio e dalle mail) applicate allo stato, con il resoconto per il report
  python3 pagamenti_aperti.py applica --state state.json --risposte risposte.json --out state.json \
      --report report.json [--voci voci.json] [--oggi AAAA-MM-GG]

risposte.json: [{"id": "C-DT-10-26-1a2b3c" | "n": 3, "risposta": "pagato" | "sollecitare" | "arrivo",
                 "data": "AAAA-MM-GG" | null, "chi": "Rossella Strati", "email": "strati@...",
                 "quando": "AAAA-MM-GG", "nota": "...", "fonte": "foglio" | "mail"}]
Con "n" (numero di riga del foglio) serve --voci per risalire all'id.
"""
import argparse
import datetime as dt
import hashlib
import json
import re
import sys

RISPOSTE = ['Pagato', 'Da sollecitare', 'In arrivo il']
BILL_ST = {'invoice': 'da fatturare', 'collect': 'da incassare', 'blocked': 'bloccato'}
AMMIN = {'strati@prestonbarbieri.com': 'Rossella Strati', 'paolabarbieri@prestonbarbieri.com': 'Paola Barbieri',
         'lorenzini@prestonbarbieri.com': 'Enrica Lorenzini'}


def oggi_di(a):
    return dt.date.fromisoformat(a.oggi) if a.oggi else dt.date.today()


def iso(s):
    return isinstance(s, str) and re.match(r'^\d{4}-\d{2}-\d{2}$', s or '') is not None


def gm(s):
    return '%s/%s/%s' % (s[8:10], s[5:7], s[0:4]) if iso(s) else ''


def voce_id(code, v):
    """Id stabile della rata: commessa + importo + inizio del testo (lo stesso criterio della pagina)."""
    k = '%d|%s' % (round(float(v.get('imp') or 0)), re.sub(r'\s+', ' ', str(v.get('c') or '')).strip()[:40])
    return 'C-%s-%s' % (code, hashlib.sha1(k.encode('utf-8')).hexdigest()[:6])


def cm_altro(c):
    g = c.get('gest') or {}
    return bool(g.get('tipo') and g.get('tipo') not in ('commessa', 'variante'))


def elenco(S, oggi):
    """Le righe del foglio: rate del contratto non incassate e scadute (o entro 30 giorni), poi le voci da fatturare."""
    lim = (oggi + dt.timedelta(days=30)).isoformat()
    out = []
    for c in S.get('commesse') or []:
        if cm_altro(c):
            continue
        for v in (c.get('pag') or {}).get('voci') or []:
            imp = float(v.get('imp') or 0)
            if v.get('inc') or v.get('st') == 'incassato':
                continue
            att = v.get('att') if iso(v.get('att')) else ''
            pa = v.get('pa') or {}
            scaduta = bool(att and att < oggi.isoformat())
            # commessa evasa o rata senza importo: entra solo se gia' scaduta (i soldi mancano comunque) o gia' in discussione
            if (c.get('ev') or not imp) and not (scaduta or pa):
                continue
            if not ((att and att <= lim) or v.get('st') in ('da_verificare', 'scaduto') or pa):
                continue
            sit = []
            if att:
                sit.append(('scaduta il ' if scaduta else 'attesa il ') + gm(att))
            if not imp:
                sit.append('importo da recuperare')
            if v.get('st') == 'da_verificare':
                sit.append('incasso da verificare')
            if c.get('sp'):
                sit.append('commessa sospesa')
            if c.get('ev'):
                sit.append('commessa evasa')
            out.append({'id': voce_id(c['code'], v), 'tipo': 'Da contratto', 'cm': c['code'], 'cli': c.get('cliente') or '',
                        'desc': re.sub(r'\s+', ' ', str(v.get('c') or '')).strip()[:180], 'imp': round(imp, 2) if imp else None, 'att': att,
                        'sit': ' · '.join(sit), 'ultima': ultima_txt(pa), 'ord': (0, att or '9999')})
    for b in S.get('billing') or []:
        if b.get('d') or b.get('st') not in BILL_ST:
            continue
        att = b.get('att') if iso(b.get('att')) else ''
        cm = ''
        m = re.search(r'\b([A-Z]{1,3}-\d{2}(?:-\d{2})?)\b', ' '.join(str(b.get(k) or '') for k in ('ride', 'w', 's')))
        if m:
            cm = m.group(1)
        out.append({'id': 'F-' + str(b.get('id')), 'tipo': 'Da fatturare / incassare', 'cm': cm, 'cli': b.get('cli') or '',
                    'desc': re.sub(r'\s+', ' ', ' — '.join(x for x in (b.get('ride'), b.get('w')) if x)).strip()[:180],
                    'imp': re.sub(r'\s+', ' ', str(b.get('a') or '')).strip()[:80], 'att': att,
                    'sit': BILL_ST[b['st']] + (' · ' + str(b.get('stx'))[:120] if b.get('stx') else ''),
                    'ultima': ultima_txt(b.get('pa') or {}), 'ord': (1, b.get('o') or '')})
    out.sort(key=lambda r: r['ord'])
    for i, r in enumerate(out, 1):
        r['n'] = i
        del r['ord']
    return out


def ultima_txt(pa):
    if not pa:
        return ''
    r = {'pagato': 'Pagato', 'sollecitare': 'Da sollecitare', 'arrivo': 'In arrivo il ' + gm(pa.get('d') or '')}.get(pa.get('r'), pa.get('r') or '')
    return '%s (%s, %s)' % (r, pa.get('chi') or '', gm(pa.get('quando') or ''))


def foglio(a):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    S = json.load(open(a.state, encoding='utf-8'))
    oggi = oggi_di(a)
    righe = elenco(S, oggi)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Pagamenti aperti'
    ws.append(['Pagamenti aperti — settimana del %s' % gm(oggi.isoformat())])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append(['Per ogni riga scegliete nella colonna RISPOSTA dal menu a tendina: Pagato, Da sollecitare oppure In arrivo il '
               '(e in quel caso scrivete la data nella colonna accanto). Note libere nell\'ultima colonna. Le risposte arrivano da sole al Second Brain di Danilo.'])
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells('A2:M2')
    ws.row_dimensions[2].height = 32
    ws.append([])
    head = ['N.', 'ID', 'Tipo', 'Commessa', 'Cliente', 'Descrizione', 'Importo €', 'Data attesa', 'Situazione nel Second Brain',
            'Ultima risposta', 'RISPOSTA', 'Data (se «In arrivo il»)', 'Note']
    ws.append(head)
    hr = ws.max_row
    fill = PatternFill('solid', fgColor='D9E1F2')
    giallo = PatternFill('solid', fgColor='FFF2CC')
    for i in range(1, len(head) + 1):
        c = ws.cell(row=hr, column=i)
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in righe:
        ws.append([r['n'], r['id'], r['tipo'], r['cm'], r['cli'], r['desc'], r['imp'],
                   dt.date.fromisoformat(r['att']) if r['att'] else None, r['sit'], r['ultima'], None, None, None])
        rr = ws.max_row
        ws.cell(row=rr, column=7).number_format = '#,##0.00'
        ws.cell(row=rr, column=8).number_format = 'DD/MM/YYYY'
        for col in (11, 12, 13):
            ws.cell(row=rr, column=col).fill = giallo
        for col in (6, 9, 10):
            ws.cell(row=rr, column=col).alignment = Alignment(wrap_text=True, vertical='top')
    last = max(ws.max_row, hr + 1)
    dv = DataValidation(type='list', formula1='"%s"' % ','.join(RISPOSTE), allow_blank=True, showDropDown=False,
                        errorTitle='Risposta non valida', error='Scegli dal menu: Pagato, Da sollecitare, In arrivo il',
                        promptTitle='Risposta', prompt='Pagato / Da sollecitare / In arrivo il (data nella colonna accanto)')
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add('K%d:K%d' % (hr + 1, last))
    dd = DataValidation(type='date', operator='greaterThan', formula1='DATE(2020,1,1)', allow_blank=True,
                        errorTitle='Data non valida', error='Scrivi una data (gg/mm/aaaa)')
    dd.showErrorMessage = True
    ws.add_data_validation(dd)
    dd.add('L%d:L%d' % (hr + 1, last))
    for col, w in zip('ABCDEFGHIJKLM', (5, 18, 16, 11, 26, 52, 13, 12, 34, 26, 16, 14, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'C%d' % (hr + 1)
    wb.save(a.out)
    json.dump({'settimana': oggi.isoformat(), 'riga_intestazione': hr, 'voci': righe}, open(a.json, 'w', encoding='utf-8'),
              ensure_ascii=False, indent=1)
    nc = sum(1 for r in righe if r['tipo'] == 'Da contratto')
    print('foglio: %d righe (%d da contratto, %d da fatturare/incassare) -> %s' % (len(righe), nc, len(righe) - nc, a.out))


def norm_risposta(x):
    t = str(x or '').strip().lower()
    if t.startswith('pagat') or t in ('p', 'incassato', 'incassata'):
        return 'pagato'
    if 'solleci' in t or t == 's':
        return 'sollecitare'
    if 'arriv' in t or t == 'a':
        return 'arrivo'
    return None


def applica(a):
    S = json.load(open(a.state, encoding='utf-8'))
    oggi = oggi_di(a)
    ris = json.load(open(a.risposte, encoding='utf-8'))
    V = {}
    if a.voci:
        V = json.load(open(a.voci, encoding='utf-8'))
    per_n = {int(r['n']): r['id'] for r in V.get('voci') or []}
    per_id = {r['id']: r for r in V.get('voci') or []}
    # la settimana del foglio in uso (senza --voci: il lunedi' di questa settimana)
    sett = V.get('settimana') or (oggi - dt.timedelta(days=oggi.weekday())).isoformat()
    # indice id → oggetto nello stato
    idx = {}
    for c in S.get('commesse') or []:
        for v in (c.get('pag') or {}).get('voci') or []:
            idx[voce_id(c['code'], v)] = ('C', c, v)
    for b in S.get('billing') or []:
        idx['F-' + str(b.get('id'))] = ('F', None, b)
    rep = {'applicate': [], 'ignorate': [], 'sollecitare': [], 'conflitti': [], 'ricondotte': []}

    def ripiego(rid):
        """Il testo della rata e' cambiato dopo il foglio (e con lui l'id): la ritrovo per commessa, importo e data attesa."""
        r = per_id.get(rid)
        if not r or r.get('tipo') != 'Da contratto':
            return None
        cand = [k for k, (t, c, v) in idx.items() if t == 'C' and c['code'] == r['cm']
                and round(float(v.get('imp') or 0)) == round(float(r.get('imp') or 0))
                and (v.get('att') or '') == (r.get('att') or '')]
        return cand[0] if len(cand) == 1 else None

    # la risposta piu' recente per id vince; risposte diverse nello stesso giorno sono un conflitto da segnalare
    migliori = {}
    for x in ris:
        n = x.get('n')
        try:
            n = int(str(n).strip()) if n not in (None, '') else None
        except ValueError:
            n = None
        if not x.get('id') and n is not None:
            # il numero di riga vale solo per il foglio di questa settimana: una risposta a un foglio precedente
            # punterebbe a un'altra rata
            if (x.get('settimana') and x.get('settimana') != sett) or (not x.get('settimana') and str(x.get('quando') or '9999') < sett):
                rep['ignorate'].append({'risposta': x, 'motivo': 'numero di riga di un foglio di una settimana precedente'})
                continue
        rid = x.get('id') or per_n.get(n)
        if rid and rid not in idx:
            alt = ripiego(rid)
            if alt:
                rep['ricondotte'].append({'da': rid, 'a': alt})
                rid = alt
        r = norm_risposta(x.get('risposta'))
        if not rid or rid not in idx or not r:
            rep['ignorate'].append({'risposta': x, 'motivo': 'riga non trovata' if (not rid or rid not in idx) else 'risposta non riconosciuta'})
            continue
        if r == 'arrivo' and not iso(x.get('data')):
            rep['ignorate'].append({'risposta': x, 'motivo': '«In arrivo il» senza una data valida'})
            continue
        email = str(x.get('email') or '').lower()
        chi = x.get('chi') or AMMIN.get(email) or email or 'amministrazione'
        y = {'id': rid, 'r': r, 'd': x.get('data') if iso(x.get('data')) else None, 'chi': chi, 'quando': x.get('quando') or oggi.isoformat(),
             'nota': str(x.get('nota') or '')[:300], 'fonte': x.get('fonte') or 'foglio'}
        p = migliori.get(rid)
        if p and p['quando'] == y['quando'] and (p['r'], p['d']) != (y['r'], y['d']):
            rep['conflitti'].append({'id': rid, 'a': p, 'b': y})
        if not p or y['quando'] >= p['quando']:
            migliori[rid] = y
    grp = None
    for g in S.get('groups') or []:
        if g.get('code') == 'PAGAMENTI':
            grp = g
    if grp is None:
        grp = {'code': 'PAGAMENTI', 'name': 'Pagamenti', 'tasks': []}
        S.setdefault('groups', []).append(grp)
    for rid, y in migliori.items():
        tipo, c, o = idx[rid]
        storia = o.get('paStoria') or ([o['pa']] if o.get('pa') else [])
        # il foglio si rilegge ogni giorno: la stessa risposta gia' registrata in questa settimana non si riapplica
        # (ne' scavalca una risposta diversa arrivata dopo per mail). La stessa risposta di una settimana precedente
        # invece e' una nuova conferma.
        if any((h.get('r'), h.get('d')) == (y['r'], y['d']) and str(h.get('quando') or '') >= sett for h in storia):
            continue
        ultima = storia[-1] if storia else {}
        if ultima and str(ultima.get('quando') or '') > y['quando']:
            rep['conflitti'].append({'id': rid, 'a': ultima, 'b': y, 'motivo': 'risposta piu\' vecchia di quella gia\' registrata: non applicata'})
            continue
        # una rata gia' incassata non torna aperta da sola: lo segnalo a Danilo
        gia = (tipo == 'C' and (o.get('inc') or o.get('st') == 'incassato')) or (tipo == 'F' and o.get('d'))
        if gia and y['r'] != 'pagato':
            rep['conflitti'].append({'id': rid, 'a': {'r': 'pagato', 'inc': o.get('inc') or o.get('dd')}, 'b': y,
                                     'motivo': 'risulta gia\' incassata: non la riapro, decide Danilo'})
            continue
        o['pa'] = {k: y[k] for k in ('r', 'd', 'chi', 'quando', 'nota', 'fonte')}
        o.setdefault('paStoria', []).append(o['pa'])
        o['paStoria'] = o['paStoria'][-10:]
        testo = {'pagato': 'pagato', 'sollecitare': 'da sollecitare', 'arrivo': 'in arrivo il ' + gm(y['d'] or '')}[y['r']]
        segno = ' · %s il %s: %s%s' % (y['chi'], gm(y['quando']), testo, (' (' + y['nota'] + ')') if y['nota'] else '')
        nome = (c['code'] + ' · ' + (c.get('cliente') or '')) if c else (o.get('cli') or '')
        desc = str(o.get('c') or o.get('w') or '')[:90]
        if tipo == 'C':
            if y['r'] == 'pagato':
                o['inc'] = y['d'] or y['quando']
                o['st'] = 'incassato'
            elif y['r'] == 'arrivo':
                o['att'] = y['d']
                if o.get('st') in ('da_verificare', 'scaduto'):
                    o['st'] = 'atteso'
            else:
                o['st'] = 'scaduto'
            o['n'] = (str(o.get('n') or '') + segno).strip(' ·')
        else:
            if y['r'] == 'pagato':
                o['d'] = True
                o['dd'] = y['quando']
            elif y['r'] == 'arrivo':
                o['att'] = y['d']
            o['stx'] = (str(o.get('stx') or '') + segno).strip(' ·')
        base = 'pgs-' + re.sub(r'[^A-Za-z0-9]+', '', rid)[-14:]
        aperti = [t for t in grp['tasks'] if (t.get('id') == base or str(t.get('id') or '').startswith(base + '-')) and not t.get('d')]
        if y['r'] in ('pagato', 'arrivo'):
            # il sollecito non serve piu': chiudo il task aperto
            for t in aperti:
                t['d'] = True
                t['dd'] = y['quando']
                t['s'] = (str(t.get('s') or '') + ' · chiuso: %s (%s, %s)' % (testo, y['chi'], gm(y['quando']))).strip(' ·')
        if y['r'] == 'sollecitare':
            usati = {t.get('id') for t in grp['tasks']}
            tid = base if base not in usati else base + '-' + y['quando'].replace('-', '')[2:]
            k = 2
            while tid in usati:
                tid = base + '-' + y['quando'].replace('-', '')[2:] + '-' + str(k)
                k += 1
            if not aperti:
                grp['tasks'].append({'id': tid, 't': 'Sollecitare %s: %s' % (nome, desc), 'o': y['quando'],
                                     's': 'Foglio pagamenti · %s · %s segnala «da sollecitare»%s' % (gm(y['quando']), y['chi'], (' · ' + y['nota']) if y['nota'] else ''),
                                     'd': False, 'n': '', 'nq': [], 'p': 1})
            rep['sollecitare'].append({'id': rid, 'chi': nome, 'desc': desc, 'da': y['chi']})
        rep['applicate'].append({'id': rid, 'tipo': 'rata del contratto' if tipo == 'C' else 'da fatturare/incassare', 'chi': nome,
                                 'desc': desc, 'risposta': testo, 'da': y['chi'], 'quando': y['quando'], 'fonte': y['fonte'], 'nota': y['nota']})
    if a.voci:
        risposte_ids = set(migliori) | {r['da'] for r in rep['ricondotte']} | {rid for rid, (t, c, o) in idx.items() if str((o.get('pa') or {}).get('quando') or '') >= sett}
        rep['senzaRisposta'] = [{'n': r['n'], 'id': r['id'], 'cm': r['cm'], 'cli': r['cli'], 'desc': r['desc'][:90], 'att': r['att']}
                                for r in V.get('voci') or [] if r['id'] not in risposte_ids]
    json.dump(S, open(a.out, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    json.dump(rep, open(a.report, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('risposte: %d applicate, %d ignorate, %d da sollecitare, %d conflitti' % (len(rep['applicate']), len(rep['ignorate']), len(rep['sollecitare']), len(rep['conflitti'])))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    f = sub.add_parser('foglio')
    f.add_argument('--state', required=True)
    f.add_argument('--out', required=True)
    f.add_argument('--json', required=True)
    f.add_argument('--oggi')
    p = sub.add_parser('applica')
    p.add_argument('--state', required=True)
    p.add_argument('--risposte', required=True)
    p.add_argument('--out', required=True)
    p.add_argument('--report', required=True)
    p.add_argument('--voci')
    p.add_argument('--oggi')
    a = ap.parse_args()
    if a.cmd == 'foglio':
        foglio(a)
    else:
        applica(a)


if __name__ == '__main__':
    sys.exit(main())
