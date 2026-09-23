#!/usr/bin/env python3
"""Budget e consuntivo per nodo → cartella Excel (.xlsx) per la scheda Commesse → Budget.

Specchio in Python del modello della pagina (funzione bdgModello del blocco R24/R25 in app.js): stessa
aritmetica, così il file Excel e la scheda mostrano gli stessi numeri. Le colonne derivate
(budget ore, budget totale, impegnato, scostamento, utile, prezzo di vendita) sono formule
Excel che leggono il foglio Parametri: Danilo può ritoccare budget, tariffe, margine e indici nel
file e vedere il ricalcolo.

R25 (18/09/2026, sera): per una commessa in corso senza budget salvato, nodi e importi vengono proposti
dalle commesse chiuse della stessa famiglia (prefisso del codice: SBC, GB, FS…; solo codici XXX-nn-aa, non
fiere, ricambi o progetti interni, non quelle che sembrano ancora in corso): per ogni nodo la somma dei
costi consuntivi indicizzati all'anno corrente divisa per il numero di commesse di riferimento (al massimo
3: prima quelle con gli stessi nodi, poi le più recenti), arrotondata a 100 €. L'utile è «a budget»
(prezzo − budget) per le commesse in corso e «consuntivo» (prezzo − costi consuntivi: impegnato + ore) per
le evase. Le ore a budget possono essere anche per l'intera commessa (c.bdg.hu / c.bdg.ho) e si sommano a
quelle per nodo. Regole allineate alla revisione finale del blocco (18/09 sera): nodo del gestionale con
lo stesso nome della commessa unito a «senza nodo»; nodo «Varie (senza nodo)» con alias '-'; colonna
Y+1 indicizzata per anno; ore delle ditte esterne nei costi consuntivi; totali delle colonne in euro come
somma dei valori arrotondati delle righe (come in tabella), ore esatte; prezzo di vendita dal contratto,
altrimenti (solo evase) dalle fatture al cliente, altrimenti dall'ordine cliente.

  budget_xlsx.py --state state.json --cm SBC-15-26 --out export/          → export/budget-SBC-15-26.xlsx
  budget_xlsx.py --state state.json --all --out export/                   → export/budget-tutte.xlsx (Riepilogo + un foglio per commessa in corso)
  budget_xlsx.py --state state.json --all --compatto --out export/        → export/budget-riepilogo.xlsx (solo Riepilogo + Parametri: piccolo, per Drive)
  budget_xlsx.py --state state.json --cm A --cm B --out file.xlsx
  budget_xlsx.py --state state.json --cm SBC-15-26 --json                 → modello in JSON (per i controlli incrociati con la pagina)
  budget_xlsx.py --state state.json --json-all                            → modello di tutte le commesse (aperte ed evase) in JSON
  budget_xlsx.py --state state.json --cm SBC-15-26 --rif SBC-14-24 …      → proposta da una sola commessa di riferimento (come il selettore in pagina)

Fogli: Riepilogo (se più commesse), un foglio per commessa (nodi, ore interne, utile), Parametri.
"""
import argparse, datetime as dt, json, math, os, re, sys

IDX_DEFAULT = {"2013": 1.2, "2014": 0.2, "2015": 0.1, "2016": -0.1, "2017": 1.2, "2018": 1.2, "2019": 0.6,
               "2020": -0.2, "2021": 1.9, "2022": 8.1, "2023": 5.7, "2024": 1.0, "2025": 1.7, "2026": 1.8, "2027": 1.9}
STADI = ('fat', 'con', 'ddt', 'ord', 'alt')
# colonne in euro: i totali sommano i valori arrotondati delle righe (BDG_MONEY in app.js)
MONEY = ('ext', 'fat', 'con', 'ddt', 'ord', 'alt', 'sost', 'imp', 'bOre', 'bTot', 'idxY', 'idxY1', 'przY', 'przY1')
_CONS_OK = {}


def num(x, d=0.0):
    """Number(x) || 0 di JavaScript."""
    if x is None or x is False or isinstance(x, (list, dict)):
        return d
    if x is True:
        return 1.0
    try:
        v = float(x)
        return d if v != v else v
    except (TypeError, ValueError):
        return d


def js_truthy(v):
    return not (v is None or v is False or v == '' or (isinstance(v, (int, float)) and (v == 0 or v != v)))


def jsround(x):
    """Math.round di JavaScript (metà verso +∞), non l'arrotondamento bancario di Python."""
    return int(math.floor(x + 0.5))


def js_keys(d):
    """Object.keys: le chiavi 'intere' prima, in ordine numerico, poi le altre in ordine di inserimento."""
    ks = list((d or {}).keys())
    ints = [k for k in ks if re.fullmatch(r'0|[1-9]\d*', str(k))]
    return sorted(ints, key=int) + [k for k in ks if k not in ints]


def bcfg(S):
    c = dict(S.get('cfg') or {})
    idx = c.get('idx') if isinstance(c.get('idx'), dict) else IDX_DEFAULT
    c['idx'] = {str(k): num(v) for k, v in idx.items()}
    c['marg'] = 30.0 if c.get('marg') is None or num(c.get('marg'), None) is None else num(c['marg'])
    c['margTipo'] = 'ricarico' if c.get('margTipo') == 'ricarico' else 'prezzo'
    c['tUff'] = num(c.get('tUff'))
    c['tOff'] = num(c.get('tOff'))
    return c


def fattore_idx(cf, da, a):
    f = 1.0
    for y in range(int(da) + 1, int(a) + 1):
        f *= 1 + cf['idx'].get(str(y), 0.0) / 100
    return f


def prezzo_vendita(cf, costo):
    m = cf['marg'] / 100
    if cf['margTipo'] == 'ricarico':
        return costo * (1 + m)
    return costo if m >= 1 else costo / (1 - m)


def norm_nodo(s):
    s = str(s or '').upper()
    for acc, rep in (('ÀÁÂÃÄ', 'A'), ('ÈÉÊË', 'E'), ('ÌÍÎÏ', 'I'), ('ÒÓÔÕÖ', 'O'), ('ÙÚÛÜ', 'U')):
        for ch in acc:
            s = s.replace(ch, rep)
    return re.sub(r'[^A-Z0-9]+', ' ', s).strip()


def titolo(s):
    return re.sub(r'(^|[\s(/-])([a-zà-ù])', lambda m: m.group(1) + m.group(2).upper(), str(s or '').lower())


def nodo_label(nd):
    return 'Senza nodo' if nd == '-' else titolo(nd)


def eur_r(n):
    """eurR della pagina: arrotondato, migliaia col punto, «€»."""
    return '—' if n is None else format(jsround(n), ',').replace(',', '.') + ' €'


def ore_csv(h):
    return jsround(h * 10) / 10


def cons_cm(S, code):
    """S.cons.cm[code]; un nodo del gestionale con lo stesso nome della commessa è «senza nodo»: viene
    unito a '-' (una volta sola, come consCm in app.js)."""
    cons = S.get('cons') or {}
    cc = (cons.get('cm') or {}).get(code)
    if not cc:
        return None
    agg = '%s|%s' % (cons.get('agg') or '', cons.get('gen') or '')
    key = (id(S), code)
    if _CONS_OK.get(key) != agg:
        _CONS_OK[key] = agg
        k, nd = norm_nodo(code), cc.get('nd')
        if nd is None:
            nd = {}
        for n in js_keys(nd):
            if n == '-' or norm_nodo(n) != k:
                continue
            dest = nd.get('-')
            if not dest:
                dest = nd['-'] = {}
            for y in js_keys(nd[n]):
                a = nd[n][y] or {}
                d = dest.get(y)
                if not d:
                    d = dest[y] = {'n': 0}
                for s in js_keys(a):
                    d[s] = num(d.get(s)) + num(a.get(s))
            del nd[n]
    return cc


def cons_nodi(S, code):
    c = cons_cm(S, code)
    return js_keys((c or {}).get('nd'))


def fam_cm(code):
    m = re.match(r'^([A-Za-z]+)-', str(code or ''))
    return m.group(1).upper() if m else None


def cm_std(code):
    return bool(re.fullmatch(r'[A-Za-z]+-\d+-\d{2}', str(code or '')))


def cm_altro(c):
    """Fiere, ricambi, progetti interni importati dal gestionale: non sono commesse."""
    g = c.get('gest') or {}
    return bool(g.get('tipo') and g['tipo'] not in ('commessa', 'variante'))


def anno_cm(S, c):
    m = re.match(r'^[A-Za-z]+-\d+-(\d{2})$', str(c.get('code') or ''))
    if m:
        return 2000 + int(m.group(1))
    cc, y = cons_cm(S, c.get('code')), 0
    for nd, anni in ((cc or {}).get('nd') or {}).items():
        for a in (anni or {}):
            if num(a) > y:
                y = int(num(a))
    return y or None


def costi_idx_nodi(S, cf, code, Y):
    """Costi consuntivi di una commessa per nodo (nome del gestionale), indicizzati all'anno Y."""
    cc, out, tot = cons_cm(S, code), {}, 0.0
    if not cc:
        return {'nodi': out, 'tot': 0.0}
    nd_all = cc.get('nd') or {}
    for nd in js_keys(nd_all):
        s = 0.0
        for y in js_keys(nd_all[nd]):
            a = nd_all[nd][y] or {}
            v = 0.0
            for k in STADI:
                v += num(a.get(k))
            s += v * fattore_idx(cf, int(num(y)), Y)
        if s:
            out[nd] = s
            tot += s
    return {'nodi': out, 'tot': tot}


def bdg_simili(S, cf, c, Y):
    """Commesse chiuse della stessa famiglia con costi nel gestionale (max 3): solo codici standard
    XXX-nn-aa (niente riparazioni, ricambi, fiere) e non quelle che sembrano ancora in corso (gest.attn).
    Prima quelle che usano gli stessi nodi della commessa (sovrapposizione ≥ 30 %), poi le più recenti,
    poi le più grandi; se qualche commessa usa i miei stessi nodi, la media si fa solo tra quelle."""
    fam = fam_cm(c.get('code'))
    if not fam:
        return []
    miei = set()
    for n in (c.get('nodi') or []):
        k = norm_nodo(n.get('n'))
        if k:
            miei.add(k)
    for nd in cons_nodi(S, c.get('code')):
        k = norm_nodo(nd)
        if k:
            miei.add(k)
    n_miei = len(miei)
    cand = []
    for x in S.get('commesse') or []:
        if x.get('code') == c.get('code') or not x.get('ev') or fam_cm(x.get('code')) != fam or not cm_std(x.get('code')) \
                or cm_altro(x) or (x.get('gest') or {}).get('attn'):
            continue
        k = costi_idx_nodi(S, cf, x.get('code'), Y)
        ov = 0
        if n_miei:
            for nd in js_keys(k['nodi']):
                if norm_nodo(nd) in miei:
                    ov += 1
        if k['tot'] > 0:
            cand.append({'c': x, 'code': x.get('code'), 'tot': k['tot'], 'nodi': k['nodi'], 'anno': anno_cm(S, x), 'ov': ov / n_miei if n_miei else 0})
    cand.sort(key=lambda s: (0 if s['ov'] >= 0.3 else 1, -(s['anno'] or 0), -s['tot']))
    aff = [s for s in cand if s['ov'] >= 0.3]
    return (aff if aff else cand)[:3]


def proposta(S, cf, c, Y, rif_mode=None):
    """Nodi e budget proposti: nodi dal contratto (c.nodi), dal gestionale della commessa e dalle commesse
    simili chiuse; importo di un nodo = somma dei costi consuntivi indicizzati di quel nodo nelle commesse
    di riferimento ÷ numero di riferimenti (un riferimento senza quel nodo conta zero), arrotondato a 100 €,
    così il totale proposto è la media dei totali. I costi senza nodo dei riferimenti vanno nel nodo
    «Varie (senza nodo)», che ha l'alias '-'. Per una commessa evasa non si propone nessun budget."""
    out, seen = [], set()
    simili = [] if c.get('ev') else bdg_simili(S, cf, c, Y)
    mode = rif_mode or 'media'
    rif = simili if mode == 'media' else [s for s in simili if s['code'] == mode]
    if not rif:
        rif, mode = simili, 'media'

    def trova(nome):
        k = norm_nodo(nome)
        if not k:
            return None
        for o in out:
            if norm_nodo(o['n']) == k or k in [norm_nodo(a) for a in o['al']]:
                return o
        for o in out:
            ko = norm_nodo(o['n'])
            if len(ko) >= 4 and len(k) >= 4 and (ko.startswith(k) or k.startswith(ko)):
                return o
        return None

    def agg(nome, d='', al=None):
        k = norm_nodo(nome)
        if not k or k in seen:
            return None
        seen.add(k)
        o = {'id': 'b%d' % (len(out) + 1), 'n': nome, 'd': d or '', 'al': list(al or []), 'ext': None, 'hu': None, 'ho': None, 'note': '', 'prop': {}}
        out.append(o)
        return o

    def varie():
        for o in out:
            if '-' in o['al']:
                return o
        v = agg('Varie (senza nodo)', 'costi registrati nel gestionale senza nodo', ['-'])
        if v and '-' not in v['al']:
            v['al'].append('-')
        return v

    def agg_gest(nd, val, ref):
        if nd == '-':
            if not (val is not None and val > 0):
                return
            hit = varie()
        else:
            hit = trova(nd)
            if hit:
                if norm_nodo(hit['n']) != norm_nodo(nd) and nd not in hit['al']:
                    hit['al'].append(nd)
            else:
                hit = agg(titolo(nd), '', [])
        if hit is None:
            return
        if val is not None and ref:
            hit['prop'][ref] = hit['prop'].get(ref, 0) + val

    for n in (c.get('nodi') or []):
        agg(n.get('n'), n.get('d'))
    for nd in cons_nodi(S, c.get('code')):
        if nd != '-':
            agg_gest(nd, None, None)
    for s in rif:
        for nd in js_keys(s['nodi']):
            agg_gest(nd, s['nodi'][nd], s['code'])
    tot = 0
    for o in out:
        refs = list(o['prop'].keys())
        if refs:
            m = 0.0
            for r in refs:
                m += o['prop'][r]
            o['ext'] = jsround(m / len(rif) / 100) * 100
            o['nrif'] = len(refs)
            tot += o['ext']
        del o['prop']
    # ore: media delle ore consuntive dei riferimenti che le hanno nel file ore
    hu = ho = None
    nO, sU, sO = 0, 0.0, 0.0
    for s in rif:
        o = s['c'].get('ore') or {}
        if o.get('uff') is not None or o.get('off') is not None:
            nO += 1
            sU += num(o.get('uff'))
            sO += num(o.get('off'))
    if nO:
        hu, ho = jsround(sU / nO), jsround(sO / nO)
    return {'nodi': out, 'rif': rif, 'simili': simili, 'mode': mode, 'tot': tot, 'hu': hu, 'ho': ho, 'nOre': nO}


def mappa(S, c, nodi):
    """Nodo del gestionale → id del nodo a budget: '-' solo tramite l'alias esplicito '-'."""
    m = {}
    norm = [{'id': b.get('id') or 'b%d' % (i + 1), 'n': norm_nodo(b.get('n')), 'al': [a for a in (norm_nodo(a) for a in (b.get('al') or [])) if a],
             'meno': '-' in (b.get('al') or [])} for i, b in enumerate(nodi)]
    for nd in cons_nodi(S, c['code']):
        if nd == '-':
            for b in norm:
                if nd not in m and b['meno']:
                    m[nd] = b['id']
            continue
        k = norm_nodo(nd)
        if not k:
            continue
        for b in norm:
            if nd in m:
                break
            if b['n'] == k or k in b['al']:
                m[nd] = b['id']
    return m


def vuota(id_, n):
    r = {'id': id_, 'n': n, 'd': '', 'al': [], 'gest': [], 'ext': 0.0, 'hu': 0.0, 'ho': 0.0, 'note': '', 'anni': {}, 'extra': False, 'bOre': 0.0, 'bTot': 0.0}
    for k in STADI:
        r[k] = 0.0
    return r


def modello(S, c, Y, rif_mode=None):
    cf = bcfg(S)
    tU, tO = cf['tUff'], cf['tOff']
    bdg = c.get('bdg') if isinstance(c.get('bdg'), dict) else {}
    salvati = js_truthy(bdg.get('nodi'))
    nodi = bdg.get('nodi') if salvati else None
    if salvati:   # tolleranza per uno stato scritto a mano: id mancante, alias come testo
        nodi = [dict(b, id=b.get('id') or 'b%d' % (i + 1), al=(b.get('al') if isinstance(b.get('al'), list) else [x.strip() for x in str(b.get('al') or '').split(',') if x.strip()]))
                for i, b in enumerate(nodi) if isinstance(b, dict)]
    prop, P = not salvati, None
    if prop:
        P = proposta(S, cf, c, Y, rif_mode)
        nodi = P['nodi']
    h_cm = {'u': num(bdg.get('hu')) if bdg.get('hu') is not None else ((P['hu'] or 0) if P else 0),
            'o': num(bdg.get('ho')) if bdg.get('ho') is not None else ((P['ho'] or 0) if P else 0)}
    cc = cons_cm(S, c['code']) or {'nd': {}, 'ric': {}}
    mp = mappa(S, c, nodi)
    righe = []
    for b in nodi:
        r = vuota(b['id'], b.get('n') or '')
        r['d'], r['al'], r['note'] = b.get('d') or '', list(b.get('al') or []), b.get('note') or ''
        r['ext'], r['hu'], r['ho'] = num(b.get('ext')), num(b.get('hu')), num(b.get('ho'))
        r['nrif'] = b.get('nrif') or 0
        righe.append(r)
    by_id = {r['id']: r for r in righe}
    extra = []
    nd_all = cc.get('nd') or {}
    for nd in js_keys(nd_all):
        anni = nd_all[nd] or {}
        r = by_id.get(mp.get(nd))
        if r is None:
            r = vuota('g:' + nd, nodo_label(nd))
            r['extra'] = True
            extra.append(r)
            by_id[r['id']] = r
        r['gest'].append(nd)
        for y in js_keys(anni):
            a, imp = anni[y] or {}, 0.0
            for k in STADI:
                v = num(a.get(k))
                r[k] += v
                imp += v
            r['anni'][y] = r['anni'].get(y, 0.0) + imp
    extra.sort(key=lambda r: (r['id'] == 'g:-', -(r['fat'] + r['con'] + r['ddt'] + r['ord'] + r['alt'])))
    tutte = righe + extra
    tot = vuota('tot', 'Totale')
    for k in ('sost', 'imp', 'idxY', 'idxY1', 'przY', 'przY1'):
        tot[k] = 0.0
    for r in tutte:
        r['sost'] = r['fat'] + r['con'] + r['ddt'] + r['alt']
        r['imp'] = r['sost'] + r['ord']
        r['bOre'] = r['hu'] * tU + r['ho'] * tO
        r['bTot'] = r['ext'] + r['bOre']
        r['sc'] = r['imp'] - r['ext']
        r['pct'] = r['imp'] / r['ext'] if r['ext'] > 0 else None
        iy = iy1 = 0.0
        for y in js_keys(r['anni']):
            iy += r['anni'][y] * fattore_idx(cf, int(num(y)), Y)
            iy1 += r['anni'][y] * fattore_idx(cf, int(num(y)), Y + 1)
        r['idxY'], r['idxY1'] = iy, iy1
        r['przY'], r['przY1'] = prezzo_vendita(cf, iy), prezzo_vendita(cf, iy1)
        if r['ext'] > 0:
            r['stato'] = 'bad' if r['pct'] > 1 else 'warn' if r['pct'] > 0.9 else 'ok'
        else:
            r['stato'] = ('extra' if r['extra'] else 'nob') if r['imp'] else 'vuoto'
        # i totali sommano i valori arrotondati che si vedono in tabella, così la colonna torna
        for k in MONEY:
            tot[k] += jsround(r[k])
        tot['hu'] += r['hu']
        tot['ho'] += r['ho']
    tot['sc'] = tot['imp'] - tot['ext']
    tot['pct'] = tot['imp'] / tot['ext'] if tot['ext'] > 0 else None
    tot['stato'] = ('bad' if tot['pct'] > 1 else 'warn' if tot['pct'] > 0.9 else 'ok') if tot['ext'] > 0 else 'nob'
    ore = c.get('ore') or {}
    cU, cO, cE, cEeur = num(ore.get('uff')), num(ore.get('off')), num(ore.get('est')), num(ore.get('estEur'))
    ore_cons = cU * tU + cO * tO + cEeur
    bU, bO = tot['hu'] + h_cm['u'], tot['ho'] + h_cm['o']
    ore_bdg = bU * tU + bO * tO
    eco = c.get('eco') if isinstance(c.get('eco'), dict) else {}
    ric = cc.get('ric') or {}
    valore = num(eco.get('valore')) if js_truthy(eco.get('valore')) else None
    oc = num(ric.get('oc')) if js_truthy(ric.get('oc')) else None
    fat_cli = num(ric.get('fat')) if js_truthy(ric.get('fat')) else None
    if valore is not None:
        prezzo, fonte = valore, 'contratto'
    elif c.get('ev') and fat_cli is not None and fat_cli > 0 and (oc is None or fat_cli >= oc):
        # R26 (20/09): per le evase vale il maggiore tra fatture al cliente e ordine cliente (le fatture nell'estrazione sono spesso parziali)
        prezzo, fonte = fat_cli, 'fatture al cliente nel gestionale'
    elif oc is not None:
        prezzo, fonte = oc, 'ordine cliente nel gestionale'
    else:
        prezzo, fonte = None, None
    ha_budget = tot['ext'] > 0 or bU > 0 or bO > 0
    M = {'c': c, 'nodi': nodi, 'proposta': prop, 'conf': bool(bdg.get('conf')), 'righe': tutte, 'tot': tot, 'Y': Y, 'tU': tU, 'tO': tO,
         'rif': P['rif'] if P else [], 'simili': P['simili'] if P else [], 'rifMode': P['mode'] if P else None, 'propTot': P['tot'] if P else 0,
         'nOre': P['nOre'] if P else 0,
         'ore': {'bU': bU, 'bO': bO, 'cU': cU, 'cO': cO, 'cE': cE, 'cEeur': cEeur, 'bEur': ore_bdg, 'cEur': ore_cons, 'hCm': h_cm},
         'prezzo': prezzo, 'prezzoFonte': fonte, 'oc': oc, 'valore': valore, 'fatCli': fat_cli, 'haBudget': ha_budget, 'chiusa': bool(c.get('ev')),
         'budgetTot': tot['ext'] + ore_bdg, 'costoOggi': tot['imp'] + ore_cons}
    # R29 (20/09): consuntivo a prezzi di quest'anno e del prossimo (esterni indicizzati, ore alle tariffe dell'anno corrente)
    M['costoIdxY'] = tot['idxY'] + ore_cons
    M['costoIdxY1'] = tot['idxY1'] + ore_cons
    # R26: un prezzo preso dal gestionale sotto la metà dei costi consuntivi è quasi certamente incompleto (fatture
    # registrate altrove): niente utile finché Danilo non scrive il prezzo; fatture e ordine cliente lontani >30% → da confermare
    M['prezzoDubbio'] = bool(c.get('ev') and prezzo is not None and fonte != 'contratto' and M['costoOggi'] > 0 and prezzo < 0.5 * M['costoOggi'])
    M['prezzoDaConfermare'] = bool(c.get('ev') and fonte != 'contratto' and prezzo is not None and not M['prezzoDubbio']
                                   and ((fat_cli is not None and fat_cli > 0 and oc is not None and abs(fat_cli - oc) / max(fat_cli, oc) > 0.3)
                                        or (fonte == 'fatture al cliente nel gestionale' and prezzo < M['costoOggi'])))
    # utile: a budget per le commesse in corso, consuntivo per le evase (R25)
    M['utileB'] = prezzo - M['budgetTot'] if prezzo is not None and ha_budget else None
    M['utileO'] = prezzo - M['costoOggi'] if prezzo is not None else None
    if c.get('ev'):
        M['utile'] = None if M['prezzoDubbio'] else M['utileO']
        M['utileTipo'] = 'consuntivo' if M['utile'] is not None else None
    else:
        M['utile'] = M['utileB']
        M['utileTipo'] = ('budget proposto' if prop and M['rif'] else 'budget') if M['utile'] is not None else None
    M['marg'] = M['utile'] / prezzo if M['utile'] is not None and prezzo else None
    M['margB'] = M['utileB'] / prezzo if M['utileB'] is not None and prezzo else None
    M['margO'] = M['utileO'] / prezzo if M['utileO'] is not None and prezzo else None
    return M


def stato_nodi(M):
    return 'evasa' if M['chiusa'] else 'proposti' if M['proposta'] else 'confermati' if M['conf'] else 'da confermare'


def rif_txt(M):
    """Descrizione della proposta: da quali commesse e come."""
    if not (M['proposta'] and M['rif']):
        return ''
    return ('media dei costi consuntivi per nodo di ' if len(M['rif']) > 1 else 'costi consuntivi per nodo di ') \
        + ', '.join(s['code'] + (' (%d)' % s['anno'] if s['anno'] else '') for s in M['rif']) + ', indicizzati al %d' % M['Y']


def prezzo_info(M):
    """Sottotitolo del prezzo come nella scheda (bdgKpi)."""
    p, fonte = M['prezzo'], M['prezzoFonte']
    t = ('da ' + fonte) if fonte else 'non trovato: né contratto né ordine cliente'
    if M['valore'] is not None and M['oc'] is not None and abs(M['valore'] - M['oc']) > 1:
        t += ' · ordine cliente ' + eur_r(M['oc'])
    if M['fatCli'] is not None and M['fatCli'] > 0 and fonte != 'fatture al cliente nel gestionale' and (p is None or abs(M['fatCli'] - p) > 1):
        t += ' · fatturato al cliente ' + eur_r(M['fatCli'])
    if fonte == 'fatture al cliente nel gestionale' and M['oc'] is not None and abs(M['oc'] - p) / p > 0.3:
        t += ' · attenzione: ordine cliente ' + eur_r(M['oc'])
    if M.get('prezzoDubbio'):
        t += ' · probabilmente incompleto: sotto la metà dei costi consuntivi, scrivi il prezzo di vendita'
    elif M.get('prezzoDaConfermare'):
        t += ' · da confermare'
    return t


def utile_label(M):
    """Etichetta dell'utile come nell'export della pagina («Utile » + tipo)."""
    if M['chiusa']:
        return 'Utile consuntivo'
    return 'Utile budget proposto' if M['proposta'] and M['rif'] else 'Utile budget'


# ---------------------------------------------------------------- Excel
def _stili():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style='thin', color='D0D5DD')
    return {
        'h1': Font(name='Calibri', size=14, bold=True), 'h2': Font(name='Calibri', size=10, color='667085'),
        'th': Font(name='Calibri', size=10, bold=True, color='FFFFFF'), 'thfill': PatternFill('solid', fgColor='1F2A44'),
        'tot': Font(name='Calibri', size=10, bold=True), 'totfill': PatternFill('solid', fgColor='EEF2F7'),
        'extra': Font(name='Calibri', size=10, italic=True, color='7A4B00'), 'bold': Font(name='Calibri', size=10, bold=True),
        'prop': Font(name='Calibri', size=10, italic=True, color='1F4E9C'),
        'norm': Font(name='Calibri', size=10), 'sub': Font(name='Calibri', size=9, color='667085'),
        'in': PatternFill('solid', fgColor='FFF8E1'), 'border': Border(top=thin, bottom=thin, left=thin, right=thin),
        'right': Alignment(horizontal='right'), 'center': Alignment(horizontal='center'), 'wrap': Alignment(wrap_text=True, vertical='top'),
        'bad': PatternFill('solid', fgColor='FDE2E1'), 'warn': PatternFill('solid', fgColor='FFF1D6'), 'ok': PatternFill('solid', fgColor='E3F5E9'),
    }


EUR = '#,##0 "€";[Red]-#,##0 "€";"–"'
ORE = '#,##0.0 "h";-#,##0.0 "h";"–"'
PCT = '+0%;-0%;0%'
PCT0 = '0%'


def nome_foglio(code, usati):
    base = re.sub(r'[\[\]\*\?/\\:]', '_', str(code))[:28] or 'commessa'
    n, i = base, 2
    while n in usati:
        n = '%s (%d)' % (base[:25], i)
        i += 1
    usati.add(n)
    return n


def sum_round(col, r1, rn):
    """Somma dei valori arrotondati all'euro (come i totali della scheda: la colonna torna con le righe)."""
    return '=SUMPRODUCT(ROUND({0}{1}:{0}{2},0))'.format(col, r1, rn)


def scrivi_parametri(wb, cf, Y, S):
    st = _stili()
    ws = wb.create_sheet('Parametri')
    ws['A1'], ws['A1'].font = 'Parametri del calcolo', st['h1']
    ws['A2'], ws['A2'].font = 'Tariffe, margine e indici come nella pagina (Denaro → Tariffe, Commesse → Budget). Le celle gialle si possono cambiare: i fogli ricalcolano.', st['h2']
    righe = [('Tariffa ufficio €/h', cf['tUff'], '#,##0.00 "€"'), ('Tariffa produzione €/h', cf['tOff'], '#,##0.00 "€"'),
             ('Margine %', cf['marg'], '0.0'), ('Tipo margine (prezzo / ricarico)', cf['margTipo'], '@'), ('Anno corrente', Y, '0')]
    for i, (k, v, f) in enumerate(righe, start=4):
        ws.cell(i, 1, k).font = st['norm']
        c = ws.cell(i, 2, v)
        c.number_format, c.fill, c.border, c.font = f, st['in'], st['border'], st['norm']
    ws.cell(9, 1, 'Fattore prezzo (costo → prezzo di vendita)').font = st['norm']
    c = ws.cell(9, 2, '=IF(B7="ricarico",1+B6/100,IF(B6>=100,1,1/(1-B6/100)))')
    c.number_format, c.font = '0.0000', st['bold']
    ws.cell(10, 1, 'margine sul prezzo: prezzo = costo ÷ (1 − margine); ricarico: prezzo = costo × (1 + margine)').font = st['sub']
    ws.cell(12, 1, 'Anno').font = st['th']; ws.cell(12, 1).fill = st['thfill']
    ws.cell(12, 2, 'Indice %').font = st['th']; ws.cell(12, 2).fill = st['thfill']
    ws.cell(12, 3, 'Fattore fino all\'anno corrente').font = st['th']; ws.cell(12, 3).fill = st['thfill']
    ws.cell(12, 4, 'Fattore fino all\'anno successivo').font = st['th']; ws.cell(12, 4).fill = st['thfill']
    anni = sorted(set(list(cf['idx'].keys()) + [str(Y), str(Y + 1)]), key=int)
    r0 = 13
    rz = r0 + len(anni) - 1
    for i, y in enumerate(anni):
        r = r0 + i
        ws.cell(r, 1, int(y)).number_format = '0'
        c = ws.cell(r, 2, cf['idx'].get(y, 0.0))
        c.number_format, c.fill, c.border = '0.0', st['in'], st['border']
        # fattore da questo anno all'anno corrente (e al successivo): prodotto (1+idx) per gli anni successivi fino a Y (Y+1)
        # (SUMPRODUCT valuta le matrici anche senza formula-matrice, in Excel come in LibreOffice)
        ws.cell(r, 3, '=EXP(SUMPRODUCT((A$%d:A$%d>A%d)*(A$%d:A$%d<=$B$8)*LN(1+B$%d:B$%d/100)))' % (r0, rz, r, r0, rz, r0, rz)).number_format = '0.0000'
        ws.cell(r, 4, '=EXP(SUMPRODUCT((A$%d:A$%d>A%d)*(A$%d:A$%d<=$B$8+1)*LN(1+B$%d:B$%d/100)))' % (r0, rz, r, r0, rz, r0, rz)).number_format = '0.0000'
    ws.cell(rz + 2, 1, 'Fonte indici: ISTAT NIC medie annue 2013-2024; 2025-2027 stime, modificabili.').font = st['sub']
    ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 14; ws.column_dimensions['C'].width = 26; ws.column_dimensions['D'].width = 26
    cons = S.get('cons') or {}
    ws.cell(rz + 4, 1, 'Consuntivo: %s%s' % (cons.get('agg') or 'non caricato', (' · ' + cons['file']) if cons.get('file') else '')).font = st['sub']
    return {'tUff': "Parametri!$B$4", 'tOff': "Parametri!$B$5", 'fp': "Parametri!$B$9", 'anni': (r0, rz)}


def scrivi_commessa(wb, S, M, P, usati):
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter as L
    st = _stili()
    c, Y = M['c'], M['Y']
    ws = wb.create_sheet(nome_foglio(c['code'], usati))
    ws['A1'] = '%s · %s%s%s' % (c['code'], c.get('cliente') or '', (' · ' + c['desc']) if c.get('desc') else '', ' · commessa evasa' if M['chiusa'] else ' · in corso')
    ws['A1'].font = st['h1']
    bdg = c.get('bdg') if isinstance(c.get('bdg'), dict) else {}
    cons = S.get('cons') or {}
    prop_rif = M['proposta'] and bool(M['rif'])
    if M['chiusa']:
        testa = 'Commessa evasa: contano i costi consuntivi (utile consuntivo = prezzo − impegnato − ore)' \
            + (' · budget salvato il %s, resta solo come memoria' % bdg['agg'] if bdg.get('agg') else ' · nessun budget salvato')
    elif prop_rif:
        testa = 'BUDGET PROPOSTO, DA CONFERMARE: %s' % rif_txt(M) \
            + ((' · ore proposte %s h ufficio, %s h produzione (%s)' % (ore_csv(M['ore']['hCm']['u']), ore_csv(M['ore']['hCm']['o']),
                ('media di %d commesse con ore nel file' % M['nOre']) if M['nOre'] > 1 else 'dall\'unica commessa con ore nel file'))
               if M['ore']['hCm']['u'] or M['ore']['hCm']['o'] else '')
    elif M['proposta']:
        testa = 'Nodi proposti (non ancora confermati) · nessuna commessa simile chiusa con costi nel gestionale%s: il budget per nodo va inserito' % ((' (famiglia %s)' % fam_cm(c['code'])) if fam_cm(c['code']) else '')
    else:
        testa = 'Budget: %s%s%s' % ('nodi confermati' if M['conf'] else 'nodi da confermare', (' · salvato il ' + bdg['agg']) if bdg.get('agg') else '', (' · ' + bdg['fonte']) if bdg.get('fonte') else '')
    ws['A2'] = '%s · consuntivo dall\'estrazione del %s · generato il %s' % (testa, cons.get('agg') or '—', dt.date.today().isoformat())
    ws['A2'].font = st['h2']

    # ---- tabella nodi (intestazioni come l'export della pagina, bdgRighe)
    H = ['Nodo', 'Nel gestionale', 'Budget esterno', 'Ore ufficio', 'Ore produzione', 'Budget ore €', 'Budget totale', 'Fatturato', 'Consegnato', 'Ordinato',
         'Impegnato', 'Scostamento €', 'Scostamento %', 'Indicizzato %d' % Y, 'Prezzo %d' % Y, 'Indicizzato %d' % (Y + 1), 'Prezzo %d' % (Y + 1), 'Note']
    r_h = 12
    for j, h in enumerate(H, start=1):
        cell = ws.cell(r_h, j, h)
        cell.font, cell.fill, cell.alignment, cell.border = st['th'], st['thfill'], st['wrap'], st['border']
    ws.row_dimensions[r_h].height = 30
    r1 = r_h + 1
    for i, r in enumerate(M['righe']):
        row = r1 + i
        nota = r['note'] or ('non a budget' if r['extra'] else '') \
            or (('proposto: ' + rif_txt(M) + ((' · in %d su %d simili' % (r['nrif'], len(M['rif']))) if r.get('nrif') and len(M['rif']) > 1 else '')) if prop_rif and r['ext'] and not r['d'] else '') or r['d']
        vals = [r['n'], ', '.join(r['gest']), r['ext'], r['hu'], r['ho'],
                '=D{0}*{1}+E{0}*{2}'.format(row, P['tUff'], P['tOff']), '=C{0}+F{0}'.format(row),
                r['fat'] + r['con'] + r['alt'], r['ddt'], r['ord'], '=H{0}+I{0}+J{0}'.format(row),
                '=IF(C{0}>0,K{0}-C{0},"")'.format(row), '=IF(C{0}>0,K{0}/C{0}-1,"")'.format(row),
                r['idxY'], '=N{0}*{1}'.format(row, P['fp']), r['idxY1'], '=P{0}*{1}'.format(row, P['fp']), nota]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row, j, v)
            cell.border = st['border']
            cell.font = st['extra'] if r['extra'] else st['norm']
            if j in (3, 4, 5, 18):
                cell.fill = st['in']
        for j in (3, 6, 7, 8, 9, 10, 11, 12, 14, 15, 16, 17):
            ws.cell(row, j).number_format = EUR
        ws.cell(row, 4).number_format = ORE; ws.cell(row, 5).number_format = ORE; ws.cell(row, 13).number_format = PCT
        ws.cell(row, 11).font = st['bold']; ws.cell(row, 15).font = st['bold']; ws.cell(row, 17).font = st['bold']
        if prop_rif and r['ext']:
            ws.cell(row, 3).font = st['prop']
    rn = r1 + len(M['righe']) - 1 if M['righe'] else r1 - 1
    rt = rn + 1
    ws.cell(rt, 1, 'Totale')
    if M['righe']:
        for j in (3, 6, 7, 8, 9, 10, 11, 14, 15, 16, 17):
            ws.cell(rt, j, sum_round(L(j), r1, rn))
        for j in (4, 5):
            ws.cell(rt, j, '=SUM({0}{1}:{0}{2})'.format(L(j), r1, rn))
        ws.cell(rt, 12, '=IF(C{0}>0,K{0}-C{0},"")'.format(rt)); ws.cell(rt, 13, '=IF(C{0}>0,K{0}/C{0}-1,"")'.format(rt))
    else:
        for j in (3, 4, 5, 6, 7, 8, 9, 10, 11, 14, 15, 16, 17):
            ws.cell(rt, j, 0)
    for j in range(1, len(H) + 1):
        cell = ws.cell(rt, j)
        cell.font, cell.fill, cell.border = st['tot'], st['totfill'], st['border']
        cell.number_format = ORE if j in (4, 5) else PCT if j == 13 else EUR if j >= 3 else '@'
    if M['righe']:
        ws.conditional_formatting.add('M%d:M%d' % (r1, rt), CellIsRule(operator='greaterThan', formula=['0'], fill=st['bad']))
        ws.conditional_formatting.add('M%d:M%d' % (r1, rt), CellIsRule(operator='between', formula=['-0.1', '0'], fill=st['warn']))
        ws.conditional_formatting.add('M%d:M%d' % (r1, rt), CellIsRule(operator='lessThan', formula=['-0.1'], fill=st['ok']))
        ws.conditional_formatting.add('L%d:L%d' % (r1, rt), CellIsRule(operator='greaterThan', formula=['0'], fill=st['bad']))

    # ---- ore interne: budget = ore per nodo (somma della tabella) + ore per l'intera commessa (cella gialla);
    #      consuntivo del gestionale per commessa, ditte esterne al costo del file ore (come bdgRighe: Ore … Note)
    ro = rt + 2
    o = M['ore']
    ws.cell(ro, 1, 'Ore interne').font = st['bold']
    ws.cell(ro, 2, 'budget = ore per nodo (somma) + ore per l\'intera commessa (celle gialle K)%s · consuntivo del gestionale per commessa%s' % (
        (' · proposte: media delle commesse di riferimento' if prop_rif and (o['hCm']['u'] or o['hCm']['o']) else ''),
        (' · ore aggiornate a ' + c['ore']['last']) if (c.get('ore') or {}).get('last') else '')).font = st['sub']
    HO = ['Ore', 'Budget h', 'Budget €', 'Consuntivo h', 'Consuntivo €', 'Note', 'Scost. €', 'Scost. %', 'Tariffa €/h', 'Ore per nodo (somma)', 'Ore intera commessa']
    for j, h in enumerate(HO, start=1):
        cell = ws.cell(ro + 1, j, h)
        cell.font, cell.fill, cell.border, cell.alignment = st['th'], st['thfill'], st['border'], st['wrap']
    ws.row_dimensions[ro + 1].height = 30
    ru, rp = ro + 2, ro + 3
    est = bool(o['cE'] or o['cEeur'])
    re_ = rp + 1 if est else None
    rs = (re_ or rp) + 1
    for row, nome, tar, col, ch, hcm in ((ru, 'Ufficio', P['tUff'], 'D', o['cU'], o['hCm']['u']), (rp, 'Produzione', P['tOff'], 'E', o['cO'], o['hCm']['o'])):
        vals = [nome, '=J{0}+K{0}'.format(row), '=B{0}*I{0}'.format(row), ch, '=D{0}*I{0}'.format(row),
                ('di cui %s h per l\'intera commessa' % ore_csv(hcm)) if hcm else '',
                '=IF(C{0}>0,E{0}-C{0},"")'.format(row), '=IF(C{0}>0,E{0}/C{0}-1,"")'.format(row), '=' + tar, '={0}{1}'.format(col, rt), hcm or 0]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row, j, v)
            cell.border, cell.font = st['border'], st['norm']
        ws.cell(row, 9).number_format = '#,##0.00 "€"'
        for j in (2, 4, 10, 11):
            ws.cell(row, j).number_format = ORE
        for j in (3, 5, 7):
            ws.cell(row, j).number_format = EUR
        ws.cell(row, 8).number_format = PCT
        ws.cell(row, 11).fill = st['in']
        if prop_rif and hcm:
            ws.cell(row, 11).font = st['prop']
    if est:
        vals = ['Ditte esterne', '', '', o['cE'], o['cEeur'], 'costo dal file ore']
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(re_, j, v)
            cell.border, cell.font = st['border'], st['norm']
        ws.cell(re_, 4).number_format = ORE; ws.cell(re_, 5).number_format = EUR
    ws.cell(rs, 1, 'Totale')
    for j in (2, 3, 10, 11):
        ws.cell(rs, j, '={0}{1}+{0}{2}'.format(L(j), ru, rp))
    for j in (4, 5):
        ws.cell(rs, j, '=SUM({0}{1}:{0}{2})'.format(L(j), ru, re_ or rp))
    ws.cell(rs, 7, '=IF(C{0}>0,E{0}-C{0},"")'.format(rs)); ws.cell(rs, 8, '=IF(C{0}>0,E{0}/C{0}-1,"")'.format(rs))
    for j in range(1, len(HO) + 1):
        cell = ws.cell(rs, j)
        cell.font, cell.fill, cell.border = st['tot'], st['totfill'], st['border']
        cell.number_format = ORE if j in (2, 4, 10, 11) else PCT if j == 8 else EUR if j in (3, 5, 7) else '@'
    ws.conditional_formatting.add('H%d:H%d' % (ru, rs), CellIsRule(operator='greaterThan', formula=['0'], fill=st['bad']))

    # ---- riquadro riepilogo (righe 4-10), con formule sulla tabella e sul blocco ore; etichette come l'export della pagina
    ws['A4'], ws['B4'] = 'Prezzo di vendita', M['prezzo'] if M['prezzo'] is not None and not M.get('prezzoDubbio') else ''
    ws['C4'] = prezzo_info(M) if not M.get('prezzoDubbio') else ('nel gestionale %s (%s), probabilmente incompleto: scrivi qui il prezzo di vendita e l\'utile si calcola' % (eur_r(M['prezzo']), M['prezzoFonte']))
    ws['A5'], ws['B5'] = 'Budget totale commessa', '=C{0}+C{1}'.format(rt, rs)
    ws['C5'] = '=IF(OR(B5>0,B{1}>0),"esterni "&FIXED(C{0},0)&" € + ore "&FIXED(C{1},0)&" €"{2},"{3}")'.format(
        rt, rs, '&" · proposto da %s, da confermare"' % ', '.join(s['code'] for s in M['rif']) if prop_rif else '', 'commessa evasa: contano i costi consuntivi' if M['chiusa'] else 'nessun budget inserito')
    ws['A6'], ws['B6'] = 'Costi consuntivi' if M['chiusa'] else 'Costi a oggi', '=K{0}+E{1}'.format(rt, rs)
    ws['C6'] = '="impegnato "&FIXED(K{0},0)&" € + ore "&FIXED(E{1},0)&" €"'.format(rt, rs)
    ws['A7'], ws['B7'] = 'Scostamento esterni', '=IF(C{0}>0,K{0}-C{0},"")'.format(rt)
    ws['C7'] = '=IF(C{0}>0,TEXT(K{0}/C{0}-1,"+0%;-0%")&" sul budget esterno{1}","{2}")'.format(
        rt, ' proposto' if prop_rif else '', 'commessa evasa: nessun budget da confrontare' if M['chiusa'] else 'serve il budget per nodo')
    utile_b = '=IF(OR(B4="",B4<=0,AND(B5<=0,B{0}<=0)),"",B4-B5)'.format(rs)   # utile a budget: solo con prezzo e budget (esterni o ore)
    utile_o = '=IF(OR(B4="",B4<=0),"",B4-B6)'                                      # utile con i costi a oggi / consuntivi
    ws['A8'] = utile_label(M)
    if M['chiusa']:
        ws['B8'] = utile_o
        ws['C8'] = '=IF(OR(B4="",B4<=0),"manca il prezzo di vendita","margine "&TEXT(B8/B4,"0%")&" sul prezzo · prezzo − costi consuntivi (impegnato + ore)")'
    else:
        ws['B8'] = utile_b
        ws['C8'] = '=IF(OR(B4="",B4<=0),"manca il prezzo di vendita",IF(B8="","serve il budget: inseriscilo o conferma quello proposto","margine "&TEXT(B8/B4,"0%")&" sul prezzo{0} · con i soli costi a oggi "&FIXED(B4-B6,0)&" € ("&TEXT((B4-B6)/B4,"0%")&")"))'.format(
            ' · sul budget proposto, da confermare' if prop_rif else '')
    ws['A9'], ws['B9'] = 'Utile a budget', utile_b
    ws['C9'] = '=IF(OR(B9="",B4<=0),"serve prezzo e budget","margine "&TEXT(B9/B4,"0%")&" · prezzo − budget totale' + (' (memoria)' if M['chiusa'] else '') + '")'
    ws['A10'], ws['B10'] = 'Utile con i costi a oggi', utile_o
    ws['C10'] = '=IF(OR(B10="",B4<=0),"manca il prezzo di vendita","margine "&TEXT(B10/B4,"0%")&" · prezzo − costi a oggi (impegnato + ore)")'
    for r in range(4, 11):
        ws.cell(r, 1).font = st['norm']; ws.cell(r, 2).font = st['bold']; ws.cell(r, 2).number_format = EUR; ws.cell(r, 3).font = st['sub']
    ws['B4'].fill = st['in']
    ws.conditional_formatting.add('B8', CellIsRule(operator='lessThan', formula=['0'], fill=st['bad']))
    ws.conditional_formatting.add('B9', CellIsRule(operator='lessThan', formula=['0'], fill=st['bad']))
    ws.conditional_formatting.add('B10', CellIsRule(operator='lessThan', formula=['0'], fill=st['bad']))
    ws.conditional_formatting.add('B7', CellIsRule(operator='greaterThan', formula=['0'], fill=st['bad']))

    # ---- larghezze, blocco, stampa
    widths = [30, 26, 14, 11, 12, 13, 14, 13, 13, 13, 14, 13, 9, 14, 14, 14, 14, 28]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = 'B%d' % r1
    ws.sheet_view.zoomScale = 90
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.cell(rs + 2, 1, 'Celle gialle = budget da compilare (per nodo; ore anche per l\'intera commessa). Fatturato = fatture fornitore + costi contabilizzati; Consegnato = DDT e conto lavoro non ancora fatturati; Ordinato = ordini fornitore aperti; Impegnato = la somma. '
            'I totali in euro sommano le righe arrotondate all\'euro, come in tabella. Le ore delle ditte esterne valgono il costo registrato nel file ore. '
            'Utile a budget = prezzo − budget totale (commesse in corso); utile consuntivo = prezzo − costi (commesse evase). Indicizzato e Prezzo usano il foglio Parametri.').font = st['sub']
    return ws.title


def scrivi_riepilogo(wb, S, MM):
    """Foglio Riepilogo: stesse colonne e valori dell'export «Copia per Excel» della vista d'insieme (bdgRigheRiep);
    come nella scheda, i progetti del gestionale che non sono commesse (fiere, ricambi, interni) restano fuori."""
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter as L
    st = _stili()
    ws = wb.create_sheet('Riepilogo', 0)
    cons = S.get('cons') or {}
    altre = [M for M in MM if cm_altro(M['c'])]
    MM = [M for M in MM if not cm_altro(M['c'])]
    n_ch = sum(1 for M in MM if M['chiusa'])
    ws['A1'], ws['A1'].font = 'Budget e consuntivo · ' + ('commesse in corso' if not n_ch else 'commesse evase' if n_ch == len(MM) else 'commesse'), st['h1']
    ws['A2'] = '%d commesse%s · consuntivo dall\'estrazione del %s · generato il %s · valori calcolati (i fogli per commessa hanno le formule) · utile: a budget per le commesse in corso, consuntivo per le evase%s' % (
        len(MM), (' (%d evase)' % n_ch) if n_ch and n_ch != len(MM) else '', cons.get('agg') or '—', dt.date.today().isoformat(),
        (' · esclusi dai totali %d progetti del gestionale che non sono commesse (fiere, ricambi, interni): hanno il loro foglio' % len(altre)) if altre else '')
    ws['A2'].font = st['h2']
    Yr = MM[0]['Y'] if MM else dt.date.today().year
    H = ['Commessa', 'Cliente', 'Descrizione', 'Nodi', 'Stato nodi', 'Prezzo di vendita', 'Fonte prezzo', 'Budget esterno', 'Budget ore €', 'Budget totale', 'Fatturato', 'Consegnato', 'Ordinato', 'Impegnato',
         'Ore consuntivo €', 'Costi a oggi', 'Scostamento esterni €', 'Scostamento %', 'Utile', 'Tipo utile', 'Margine %', 'Utile a budget', 'Utile con i costi a oggi',
         'Costi a prezzi %d' % Yr, 'Costi a prezzi %d' % (Yr + 1)]
    rh = 4
    for j, h in enumerate(H, start=1):
        cell = ws.cell(rh, j, h)
        cell.font, cell.fill, cell.alignment, cell.border = st['th'], st['thfill'], st['wrap'], st['border']
    ws.row_dimensions[rh].height = 30
    for i, M in enumerate(MM):
        row, t, c = rh + 1 + i, M['tot'], M['c']
        vals = [c['code'], c.get('cliente') or '', c.get('desc') or '', len(M['nodi']), stato_nodi(M),
                M['prezzo'] if M['prezzo'] is not None and not M.get('prezzoDubbio') else '',
                ('incompleto nel gestionale: %s da %s' % (eur_r(M['prezzo']), (M['prezzoFonte'] or '').replace(' nel gestionale', ''))) if M.get('prezzoDubbio') else (M['prezzoFonte'] or ''),
                t['ext'], M['ore']['bEur'], M['budgetTot'], t['fat'] + t['con'] + t['alt'], t['ddt'], t['ord'], t['imp'],
                M['ore']['cEur'], M['costoOggi'], t['sc'] if t['ext'] > 0 else '', (t['pct'] - 1) if t['pct'] is not None else '',
                M['utile'] if M['utile'] is not None else '', M['utileTipo'] or ('prezzo incompleto' if M.get('prezzoDubbio') else ''), M['marg'] if M['marg'] is not None else '',
                M['utileB'] if M['utileB'] is not None else '', M['utileO'] if M['utileO'] is not None else '',
                M['costoIdxY'], M['costoIdxY1']]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row, j, v)
            cell.border, cell.font = st['border'], st['norm']
            cell.number_format = PCT if j == 18 else PCT0 if j == 21 else '@' if j in (7, 20) else EUR if j >= 6 else '@' if j != 4 else '0'
        ws.cell(row, 1).font = st['bold']; ws.cell(row, 19).font = st['bold']
        if not M['chiusa'] and (M['proposta'] or not M['conf']):
            ws.cell(row, 5).fill = st['warn']
    rn = rh + len(MM)
    rt = rn + 1
    ws.cell(rt, 1, 'Totale')
    for j in (6, 8, 9, 10, 11, 12, 13, 14, 15, 16, 19, 22, 23, 24, 25):
        ws.cell(rt, j, '=SUM({0}{1}:{0}{2})'.format(L(j), rh + 1, rn)).number_format = EUR
    ws.cell(rt, 17, '=IF(H{0}>0,N{0}-H{0},"")'.format(rt)).number_format = EUR
    ws.cell(rt, 18, '=IF(H{0}>0,N{0}/H{0}-1,"")'.format(rt)).number_format = PCT
    ws.cell(rt, 20, 'consuntivo' if n_ch == len(MM) and MM else 'a budget / consuntivo' if n_ch else 'a budget')
    ws.cell(rt, 21, '=IF(F{0}>0,S{0}/F{0},"")'.format(rt)).number_format = PCT0
    for j in range(1, len(H) + 1):
        cell = ws.cell(rt, j)
        cell.font, cell.fill, cell.border = st['tot'], st['totfill'], st['border']
    if MM:
        ws.conditional_formatting.add('R%d:R%d' % (rh + 1, rt), CellIsRule(operator='greaterThan', formula=['0'], fill=st['bad']))
        ws.conditional_formatting.add('S%d:S%d' % (rh + 1, rt), CellIsRule(operator='lessThan', formula=['0'], fill=st['bad']))
    widths = [12, 26, 24, 6, 18, 15, 22, 14, 13, 14, 13, 13, 13, 14, 14, 14, 15, 9, 14, 15, 10, 14, 16]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = 'B%d' % (rh + 1)
    ws.auto_filter.ref = 'A%d:%s%d' % (rh, L(len(H)), max(rn, rh + 1))
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def lista_commesse(S, codici):
    """--all: le commesse in corso come le vede la scheda (non evase, non sospese, non fiere/ricambi/interni)."""
    by = {c.get('code'): c for c in S.get('commesse', [])}
    if codici is None:
        return [c for c in S.get('commesse', []) if not c.get('ev') and not c.get('sp') and not cm_altro(c)]
    mancanti = [k for k in codici if k not in by]
    if mancanti:
        sys.exit('commesse non trovate nello stato: %s' % ', '.join(mancanti))
    return [by[k] for k in codici]


def costruisci(S, codici, Y=None, riepilogo=None, rif_mode=None, compatto=False):
    """compatto=True: solo Riepilogo + Parametri (niente fogli per commessa): file piccolo, adatto al caricamento automatico su Drive."""
    import openpyxl
    Y = Y or dt.date.today().year
    cf = bcfg(S)
    lista = lista_commesse(S, codici)
    MM = [modello(S, c, Y, rif_mode) for c in lista]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    P = scrivi_parametri(wb, cf, Y, S)
    usati = {'Parametri', 'Riepilogo'}
    if not compatto:
        for M in MM:
            scrivi_commessa(wb, S, M, P, usati)
    if compatto or riepilogo or (riepilogo is None and len(MM) != 1):
        scrivi_riepilogo(wb, S, MM)
    wb.move_sheet('Parametri', offset=len(wb.sheetnames))
    wb.active = 0
    return wb, MM


def json_modello(MM):
    def sim(s):
        return {'code': s['code'], 'anno': s['anno'], 'tot': s['tot'], 'nodi': s['nodi'], 'ov': s.get('ov', 0)}

    def pulisci(M):
        d = {k: v for k, v in M.items() if k not in ('c', 'righe', 'tot', 'nodi', 'rif', 'simili')}
        d['code'] = M['c']['code']
        d['nodi'] = [{k: v for k, v in n.items()} for n in M['nodi']]
        d['rif'] = [sim(s) for s in M['rif']]
        d['simili'] = [sim(s) for s in M['simili']]
        d['righe'] = [{k: v for k, v in r.items() if k != 'anni'} | {'anni': r['anni']} for r in M['righe']]
        d['tot'] = M['tot']
        return d
    return [pulisci(M) for M in MM]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--state', required=True)
    ap.add_argument('--cm', action='append', help='codice commessa (ripetibile)')
    ap.add_argument('--all', action='store_true', help='tutte le commesse in corso (non evase, non sospese, non fiere/ricambi/interni)')
    ap.add_argument('--out', help='file .xlsx o cartella di destinazione')
    ap.add_argument('--anno', type=int, help='anno corrente (default: oggi)')
    ap.add_argument('--rif', help='commessa chiusa di riferimento per la proposta (default: media delle simili)')
    ap.add_argument('--json', action='store_true', help='stampa il modello calcolato in JSON invece di scrivere il file')
    ap.add_argument('--json-all', action='store_true', help='modello in JSON di tutte le commesse dello stato (aperte ed evase), per i controlli incrociati')
    ap.add_argument('--compatto', action='store_true', help='solo Riepilogo e Parametri, senza i fogli per commessa (file piccolo per Drive)')
    a = ap.parse_args()
    if not a.all and not a.cm and not a.json_all:
        ap.error('indica --cm CODICE (anche più volte), --all oppure --json-all')
    S = json.load(open(a.state, encoding='utf-8'))
    codici = None if a.all else a.cm
    if a.json or a.json_all:
        Y = a.anno or dt.date.today().year
        lista = list(S.get('commesse') or []) if a.json_all else lista_commesse(S, codici)
        print(json.dumps(json_modello([modello(S, c, Y, a.rif) for c in lista]), ensure_ascii=False, indent=1))
        return
    if not a.out:
        ap.error('--out è obbligatorio senza --json')
    wb, MM = costruisci(S, codici, a.anno, rif_mode=a.rif, compatto=a.compatto)
    if not MM:
        sys.exit('nessuna commessa da esportare' + (' (nessuna commessa in corso nello stato)' if codici is None else ''))
    out = a.out
    if os.path.isdir(out) or out.endswith('/') or not out.lower().endswith('.xlsx'):
        os.makedirs(out, exist_ok=True)
        nome = ('budget-riepilogo.xlsx' if a.compatto else 'budget-tutte.xlsx') if codici is None else 'budget-%s.xlsx' % re.sub(r'[^A-Za-z0-9_-]+', '_', '-'.join(codici))[:60]
        out = os.path.join(out, nome)
    wb.save(out)
    print('%s · %d commesse · fogli: %s' % (out, len(MM), ', '.join(wb.sheetnames[:6]) + (' …' if len(wb.sheetnames) > 6 else '')))


if __name__ == '__main__':
    main()
