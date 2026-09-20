#!/usr/bin/env python3
"""Resoconto dei punti aperti su tutte le commesse (Commesse → Budget) → cartella Excel.

Richiesta di Danilo del 20/09/2026: «Fai tutte le commesse del file del gestionale. Fammi un resoconto dei punti
aperti». Legge state.json (commesse, cons, cfg) e usa il modello di budget_xlsx.py (specchio della scheda) per
elencare, commessa per commessa, cosa manca o non torna:

  In corso     prezzo, budget proposto (da quali commesse chiuse), costi a oggi, utile a budget, segnalazioni, dati mancanti
  Evase        prezzo trovato nel gestionale (fatture al cliente / ordine cliente), costi consuntivi, utile, segnalazioni
  Senza prezzo evase di tipo commessa senza prezzo o con prezzo incompleto: colonna vuota da compilare
  Controlli    possibili doppioni e consegne parziali segnalati da consuntivo.py, per commessa
  Riepilogo    numeri e cose da fare

  punti_aperti.py --state state.json --out export/            → export/punti-aperti-AAAA-MM-GG.xlsx
"""
import argparse, datetime as dt, json, os, sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import budget_xlsx as B  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit('serve openpyxl: pip install openpyxl')

EUR = '#,##0 €'
PCT = '0%'


def tipo_cm(c):
    g = c.get('gest') if isinstance(c.get('gest'), dict) else {}
    return g.get('tipo') or ('manuale' if not g else 'commessa')


def k_eur(x):
    return B.eur_r(x) if x is not None else '—'


def segnala_aperta(c, M):
    s, manca = [], []
    if M['prezzo'] is None:
        s.append('senza prezzo di vendita')
    if c.get('sp'):
        s.append('sospesa')
    if M['proposta'] and M['rif']:
        s.append('nodi e budget proposti da %s: da confermare in pagina' % ', '.join(r['code'] for r in M['rif']))
    elif M['proposta']:
        s.append('nessuna commessa chiusa simile: budget da inserire a mano')
    elif not M['conf']:
        s.append('budget inserito ma nodi non confermati')
    if M['prezzo'] and M['propTot'] and M['propTot'] > M['prezzo']:
        s.append('budget proposto sopra il prezzo')
    if M['prezzo'] and M['costoOggi'] > M['prezzo']:
        s.append('costi a oggi sopra il prezzo')
    elif M['haBudget'] and M['costoOggi'] > M['budgetTot']:
        s.append('costi a oggi sopra il budget%s' % (' proposto' if M['proposta'] and M['rif'] else ''))
    if M['oc'] is not None and M['valore'] is not None and abs(M['oc'] - M['valore']) > 1:
        s.append('ordine cliente nel gestionale %s diverso dal contratto' % B.eur_r(M['oc']))
    if M['margB'] is not None and 0 <= M['margB'] < 0.3:
        s.append('margine a budget sotto il 30%%: %d%%' % round(M['margB'] * 100))
    eco = c.get('eco') if isinstance(c.get('eco'), dict) else {}
    if not eco.get('valore'):
        manca.append('valore a contratto')
    if not c.get('capo'):
        manca.append('capo commessa')
    if not c.get('dc') and not c.get('scad'):
        manca.append('data di consegna')
    ore = c.get('ore') if isinstance(c.get('ore'), dict) else {}
    if not ore.get('uff') and not ore.get('off'):
        manca.append('ore nel file ore')
    return s, manca


def segnala_evasa(c, M, t):
    s = []
    if t not in ('commessa', 'variante'):
        s.append('non è una commessa (%s): fuori dai totali' % t)
    if M['prezzo'] is None:
        s.append('senza prezzo: né contratto, né fatture al cliente, né ordine cliente')
    elif M.get('prezzoDubbio'):
        s.append('prezzo del gestionale incompleto (%s, %s, costi %s): scrivi il prezzo di vendita' % (B.eur_r(M['prezzo']), M['prezzoFonte'].replace(' nel gestionale', ''), B.eur_r(M['costoOggi'])))
    elif M.get('prezzoDaConfermare'):
        s.append('fatture al cliente %s e ordine cliente %s lontani: prezzo da confermare' % (B.eur_r(M['fatCli']), B.eur_r(M['oc'])))
    if M['utile'] is not None and M['utile'] < 0:
        s.append('utile negativo: %s' % B.eur_r(M['utile']))
    elif M['marg'] is not None and M['marg'] < 0.10:
        s.append('margine sotto il 10%%: %d%%' % round(M['marg'] * 100))
    g = c.get('gest') if isinstance(c.get('gest'), dict) else {}
    if g.get('attn'):
        s.append('sembra ancora in corso: %s' % g['attn'])
    if M['conf'] is False and M['haBudget']:
        pass
    return s


def stili():
    return {'h1': Font(bold=True, size=14), 'h2': Font(italic=True, size=9, color='555555'), 'th': Font(bold=True, color='FFFFFF'),
            'thfill': PatternFill('solid', fgColor='1F3A5F'), 'bold': Font(bold=True), 'warn': PatternFill('solid', fgColor='FFF3C4'),
            'bad': PatternFill('solid', fgColor='F8D7DA'), 'wrap': Alignment(wrap_text=True, vertical='top')}


def tabella(ws, r0, H, righe, formati, st, larghezze=None, evidenzia=None):
    for j, h in enumerate(H, start=1):
        cell = ws.cell(r0, j, h)
        cell.font, cell.fill, cell.alignment = st['th'], st['thfill'], st['wrap']
    ws.row_dimensions[r0].height = 30
    for i, vals in enumerate(righe):
        r = r0 + 1 + i
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(r, j, v if v is not None else '')
            f = formati.get(j)
            if f:
                cell.number_format = f
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        if evidenzia:
            fill = evidenzia(vals)
            if fill:
                ws.cell(r, 1).fill = st[fill]
    for j, w in enumerate(larghezze or [], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    ws.auto_filter.ref = '%s%d:%s%d' % ('A', r0, get_column_letter(len(H)), r0 + max(len(righe), 1))
    return r0 + 1 + len(righe)


def costruisci(S, Y=None):
    Y = Y or dt.date.today().year
    cons = S.get('cons') or {}
    st = stili()
    wb = Workbook()
    cms = list(S.get('commesse') or [])
    MM = {c['code']: B.modello(S, c, Y) for c in cms}
    aperte = [c for c in cms if not c.get('ev')]
    evase = [c for c in cms if c.get('ev')]

    # ---- In corso
    ws = wb.active
    ws.title = 'In corso'
    ws['A1'] = 'Commesse in corso: prezzo, budget proposto, costi a oggi e cosa manca'
    ws['A1'].font = st['h1']
    ws['A2'] = "budget «proposto» = media per nodo dei costi consuntivi delle commesse chiuse simili, indicizzata al %d, da confermare in pagina (Commesse → Budget) · costi a oggi = esterni impegnati (fatturato + contabilizzato + consegnato + ordinato) + ore" % Y
    ws['A2'].font = st['h2']
    H = ['Commessa', 'Cliente', 'Descrizione', 'Stato', 'Capo', 'Prezzo di vendita', 'Fonte prezzo', 'Budget', 'Stato budget', 'Commesse di riferimento',
         'Costi a oggi', 'Utile a budget', 'Margine a budget', 'Nodi nel gestionale', 'Segnalazioni', 'Dati mancanti']
    righe, n_conf, n_prop, n_sp, n_nop = [], 0, 0, 0, 0
    for c in aperte:
        M = MM[c['code']]
        seg, manca = segnala_aperta(c, M)
        stato_b = ('proposto, da confermare' if M['proposta'] and M['rif'] else 'nessuno: da inserire' if M['proposta'] else 'confermato' if M['conf'] else 'inserito, nodi da confermare')
        if M['proposta'] and M['rif']:
            n_prop += 1
        elif M['proposta']:
            n_nop += 1
        elif M['conf']:
            n_conf += 1
        if c.get('sp'):
            n_sp += 1
        nd = len(B.cons_nodi(S, c['code']))
        righe.append([c['code'], c.get('cliente') or '', c.get('desc') or '', 'sospesa' if c.get('sp') else 'in corso', c.get('capo') or '',
                      M['prezzo'], (M['prezzoFonte'] or '').replace(' nel gestionale', ''), M['budgetTot'] if M['haBudget'] else None, stato_b,
                      ', '.join('%s (%s)' % (r['code'], r['anno']) for r in M['rif']), M['costoOggi'], M['utileB'], M['margB'], nd,
                      '\n'.join(seg), ', '.join(manca)])
    tabella(ws, 4, H, righe, {6: EUR, 8: EUR, 11: EUR, 12: EUR, 13: PCT}, st, [11, 26, 30, 9, 12, 14, 12, 14, 20, 30, 14, 14, 9, 8, 60, 34],
            evidenzia=lambda v: 'bad' if (v[5] is None or (v[11] is not None and v[11] < 0)) else ('warn' if 'confermare' in (v[8] or '') or 'inserire' in (v[8] or '') else None))

    # ---- Evase
    ws = wb.create_sheet('Evase')
    ws['A1'] = 'Commesse evase: prezzo trovato nel gestionale, costi consuntivi, utile'
    ws['A1'].font = st['h1']
    ws['A2'] = 'prezzo = valore a contratto se noto, altrimenti il maggiore tra fatture al cliente e ordine cliente nel gestionale · utile = prezzo − costi consuntivi (impegnato + ore) · un prezzo sotto la metà dei costi è considerato incompleto e non dà utile'
    ws['A2'].font = st['h2']
    H = ['Commessa', 'Tipo', 'Cliente', 'Descrizione', 'Evasa il', 'Prezzo di vendita', 'Fonte prezzo', 'Fatture al cliente', 'Ordine cliente', 'Costi consuntivi', 'Utile', 'Margine', 'Segnalazioni']
    righe, conta = [], defaultdict(int)
    for c in sorted(evase, key=lambda c: (c.get('evd') or ''), reverse=True):
        M = MM[c['code']]
        t = tipo_cm(c)
        seg = segnala_evasa(c, M, t)
        for x in seg:
            conta[x.split(':')[0].split(' (')[0]] += 1
        righe.append([c['code'], t, c.get('cliente') or '', c.get('desc') or '', c.get('evd') or '', M['prezzo'], (M['prezzoFonte'] or '').replace(' nel gestionale', ''),
                      M['fatCli'], M['oc'], M['costoOggi'], M['utile'], M['marg'], '\n'.join(seg)])
    tabella(ws, 4, H, righe, {6: EUR, 8: EUR, 9: EUR, 10: EUR, 11: EUR, 12: PCT}, st, [11, 10, 26, 30, 11, 14, 12, 14, 14, 14, 14, 9, 70],
            evidenzia=lambda v: 'bad' if (v[10] is not None and v[10] < 0) or 'incompleto' in (v[12] or '') else ('warn' if v[12] else None))

    # ---- Senza prezzo (evase di tipo commessa/variante senza prezzo o con prezzo incompleto)
    ws = wb.create_sheet('Senza prezzo')
    ws['A1'] = 'Evase senza un prezzo di vendita attendibile: compila la colonna «Prezzo di vendita (scrivi qui)»'
    ws['A1'].font = st['h1']
    ws['A2'] = 'ordinate per costi consuntivi decrescenti · lo stesso prezzo si può scrivere in pagina, editor del budget della commessa (scheda Commesse → Budget)'
    ws['A2'].font = st['h2']
    H = ['Commessa', 'Tipo', 'Cliente', 'Descrizione', 'Evasa il', 'Costi consuntivi', 'Nel gestionale', 'Prezzo di vendita (scrivi qui)']
    righe = []
    for c in evase:
        M = MM[c['code']]
        t = tipo_cm(c)
        if t not in ('commessa', 'variante'):
            continue
        if M['prezzo'] is not None and not M.get('prezzoDubbio'):
            continue
        nel = '' if M['prezzo'] is None else '%s da %s (incompleto)' % (B.eur_r(M['prezzo']), (M['prezzoFonte'] or '').replace(' nel gestionale', ''))
        righe.append([c['code'], t, c.get('cliente') or '', c.get('desc') or '', c.get('evd') or '', M['costoOggi'], nel, None])
    righe.sort(key=lambda v: -(v[5] or 0))
    tabella(ws, 4, H, righe, {6: EUR, 8: EUR}, st, [11, 10, 26, 30, 11, 14, 34, 22])
    n_senza = len(righe)

    # ---- Controlli del consuntivo per commessa
    ws = wb.create_sheet('Controlli gestionale')
    ws['A1'] = 'Segnalazioni di consuntivo.py sull\'estrazione del gestionale (i numeri del consuntivo non sono stati toccati)'
    ws['A1'].font = st['h1']
    ws['A2'] = ' · '.join(cons.get('note') or []) or 'nessuna nota'
    ws['A2'].font = st['h2']
    ctrl = cons.get('controlli') or {}
    nomi = {'ddt_gemelli': 'DDT identici a righe di fattura €', 'ordini_con_consegne': 'Ordini aperti con consegne/fatture dello stesso articolo €', 'contabilizzati_gemelli': 'Costi contabilizzati con fattura uguale nello stesso mese €'}
    chiavi = [k for k in ('ddt_gemelli', 'ordini_con_consegne', 'contabilizzati_gemelli') if k in ctrl]
    per_cm = defaultdict(dict)
    for k in chiavi:
        for code, imp in (ctrl[k].get('commesse') or {}).items():
            per_cm[code][k] = imp
    H = ['Commessa', 'Stato', 'Cliente'] + [nomi[k] for k in chiavi] + ['Totale segnalato €']
    righe = []
    for code, d in per_cm.items():
        c = next((x for x in cms if x['code'] == code), None)
        righe.append([code, ('evasa' if c.get('ev') else 'in corso') if c else '?', (c or {}).get('cliente') or ''] + [d.get(k) for k in chiavi] + [sum(d.values())])
    righe.sort(key=lambda v: -(v[-1] or 0))
    tabella(ws, 4, H, righe, {j: EUR for j in range(4, 5 + len(chiavi))}, st, [12, 9, 26] + [26] * len(chiavi) + [16])

    # ---- Riepilogo (primo foglio)
    ws = wb.create_sheet('Riepilogo', 0)
    ws['A1'] = 'Punti aperti sulle commesse · %s' % dt.date.today().strftime('%d/%m/%Y')
    ws['A1'].font = st['h1']
    ws['A2'] = 'consuntivo dall\'estrazione del gestionale del %s · %d commesse nello stato (%d in corso, %d evase)' % (cons.get('agg') or '—', len(cms), len(aperte), len(evase))
    ws['A2'].font = st['h2']
    ev_std = [c for c in evase if tipo_cm(c) in ('commessa', 'variante')]
    ev_con_utile = [c for c in ev_std if MM[c['code']]['utile'] is not None]
    ev_neg = [c for c in ev_con_utile if MM[c['code']]['utile'] < 0]
    ev_dubbio = [c for c in ev_std if MM[c['code']].get('prezzoDubbio')]
    ev_conf = [c for c in ev_std if MM[c['code']].get('prezzoDaConfermare') and not MM[c['code']].get('prezzoDubbio')]
    ev_attn = [c for c in evase if (c.get('gest') or {}).get('attn')]
    tP = sum(MM[c['code']]['prezzo'] for c in ev_con_utile)
    tU = sum(MM[c['code']]['utile'] for c in ev_con_utile)
    righe = [
        ('Commesse in corso', len(aperte), '%d sospese' % n_sp),
        ('  con budget proposto dalle commesse chiuse simili, da confermare', n_prop, 'apri la commessa in Commesse → Budget e premi «Salva e confermo i nodi» (o correggi gli importi)'),
        ('  senza commesse simili chiuse: budget da scrivere a mano', n_nop, ''),
        ('  con budget confermato', n_conf, ''),
        ('  senza prezzo di vendita', sum(1 for c in aperte if MM[c['code']]['prezzo'] is None), 'scrivilo nell\'editor del budget (campo «Prezzo di vendita»)'),
        ('Commesse evase (tipo commessa o variante)', len(ev_std), '%d fiere, ricambi, interni e altro fuori dai totali' % (len(evase) - len(ev_std))),
        ('  con prezzo attendibile e utile consuntivo', len(ev_con_utile), 'prezzo totale %s · utile %s · margine %d%%' % (B.eur_r(tP), B.eur_r(tU), round(100 * tU / tP) if tP else 0)),
        ('  con utile consuntivo negativo', len(ev_neg), 'foglio Evase: controlla prezzo e costi'),
        ('  senza prezzo o con prezzo incompleto nel gestionale', n_senza, 'foglio «Senza prezzo»: %d senza nulla, %d con fatture al cliente sotto la metà dei costi' % (n_senza - len(ev_dubbio), len(ev_dubbio))),
        ('  con fatture al cliente e ordine cliente lontani (prezzo da confermare)', len(ev_conf), 'foglio Evase, colonna Segnalazioni'),
        ('  che sembrano ancora in corso', len(ev_attn), ', '.join(c['code'] for c in ev_attn)),
    ]
    r = 4
    ws.cell(r, 1, 'Voce').font = st['th']; ws.cell(r, 2, 'Quante').font = st['th']; ws.cell(r, 3, 'Cosa fare / dettaglio').font = st['th']
    for j in (1, 2, 3):
        ws.cell(r, j).fill = st['thfill']
    for i, (a, b, c_) in enumerate(righe, start=1):
        ws.cell(r + i, 1, a); ws.cell(r + i, 2, b); ws.cell(r + i, 3, c_)
        ws.cell(r + i, 3).alignment = Alignment(wrap_text=True, vertical='top')
        if not a.startswith('  '):
            ws.cell(r + i, 1).font = st['bold']
    r += len(righe) + 2
    ws.cell(r, 1, 'Controlli sull\'estrazione del gestionale').font = st['bold']
    for i, nota in enumerate(cons.get('note') or [], start=1):
        ws.cell(r + i, 1, nota)
    ws.column_dimensions['A'].width = 70; ws.column_dimensions['B'].width = 9; ws.column_dimensions['C'].width = 90
    return wb, {'aperte': len(aperte), 'prop': n_prop, 'nop': n_nop, 'conf': n_conf, 'sp': n_sp, 'evase': len(evase), 'ev_std': len(ev_std), 'ev_utile': len(ev_con_utile),
                'ev_neg': len(ev_neg), 'senza': n_senza, 'dubbio': len(ev_dubbio), 'da_conf': len(ev_conf), 'attn': [c['code'] for c in ev_attn], 'tP': tP, 'tU': tU}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--state', required=True)
    ap.add_argument('--out', required=True, help='file .xlsx o cartella')
    ap.add_argument('--anno', type=int)
    ap.add_argument('--json', action='store_true', help='stampa i conteggi in JSON')
    a = ap.parse_args()
    S = json.load(open(a.state, encoding='utf-8'))
    wb, n = costruisci(S, a.anno)
    out = a.out
    if os.path.isdir(out) or out.endswith('/') or not out.lower().endswith('.xlsx'):
        os.makedirs(out, exist_ok=True)
        out = os.path.join(out, 'punti-aperti-%s.xlsx' % dt.date.today().isoformat())
    wb.save(out)
    if a.json:
        print(json.dumps(n, ensure_ascii=False))
    else:
        print('%s · %s' % (out, ' · '.join('%s %s' % (k, v) for k, v in n.items())))


if __name__ == '__main__':
    main()
