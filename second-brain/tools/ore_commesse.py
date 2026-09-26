#!/usr/bin/env python3
"""Ore consuntivate di TUTTE le commesse (aperte, sospese, evase) dal foglio «Ore» del gestionale (R33, 25/09/2026).

Richiesta di Danilo del 25/09: le ore andavano solo sulle commesse in corso (39 su 371); le commesse chiuse restavano
senza ore, e senza ore non c'e' ne' il costo consuntivo completo ne' la proposta di ore per le commesse nuove.

Regole (le stesse del caricamento dell'11/09, validate su 9 commesse identiche in tutti i campi):
  uff    = ore dei reparti UFF e TEC con ditta P&B
  off    = ore dei reparti OFF e ART con ditta P&B
  est    = ore delle righe con ditta diversa da P&B (Ex, ART, BBC, ...); estEur = la loro valorizzazione
  estSU / estSO = le ore esterne SENZA valorizzazione (cella vuota o «#ERRORE»: nell'estrazione dell'11/09 sono ~17.000 h
           di ex dipendenti in UFF), divise fra ufficio (UFF/TEC) e produzione (OFF/ART): la pagina le costa alle
           tariffe orarie, altrimenti varrebbero zero
  valG   = valorizzazione totale del gestionale (le righe UFF non ne hanno)
  storia = ore totali per mese [AAAA-MM, ore]; dip / att / nd = dipendenti, attivita' e nodi con piu' ore

  python3 ore_commesse.py --xlsx estrazione.xlsx --state state.json [--out state.json] [--file NOME] [--dry]
  python3 ore_commesse.py --csv ore.csv ...     (la scheda «Ore» del foglio Google preparato su Drive, separatore , o ;)
  python3 ore_commesse.py --pkl ore.pkl ...     ({'hdr': [...], 'rows': [[...], ...]})

Codici diversi: se Luca indica che le ore di una commessa stanno sotto un altro codice progetto, si scrive nello stato
oreG.alias = {"PS-16-14": ["CODICE-NEL-FOGLIO-ORE"], ...}: lo script somma quei progetti alla commessa a ogni giro.
Un alias che punta a un'altra commessa della pagina, o a un progetto gia' sommato a un'altra commessa, viene saltato
(le ore sarebbero contate due volte) e finisce in oreG.controllo.aliasSaltati.

Un'estrazione con registrazioni piu' vecchie dell'ultimo caricamento (oreG.agg) viene rifiutata: non si torna indietro
(--forza per farlo apposta).

Un foglio «Ore» rotto (zero ore totali, un solo reparto, valorizzazione tutta in errore: e' successo con l'estrazione
del 22/09) viene rifiutato e lo stato non si tocca.
"""
import argparse
import copy
import csv
import datetime as dt
import io
import json
import pickle
import re
import sys
from collections import defaultdict

COLONNE = ['Anno', 'Mese', 'Giorno', 'Progetto', 'Nodo', 'Descrizione attività', 'Dipendente', 'Reparto', 'Ditta',
           'Attività', 'Ore (pivot)', 'Valorizzazione']
REP_UFF = {'UFF', 'TEC'}
REP_OFF = {'OFF', 'ART'}


def norm_code(p):
    p = str(p or '').strip().upper()
    return '' if p in ('MANCANTE', '#N/A', 'N/A', 'NONE') else p


def num(x):
    if x is None or x == '':
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace('€', '').replace(' ', '')
    if not s or s.startswith('#'):
        return 0.0
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


def pulito(s):
    return re.sub(r'\s+', ' ', str(s or '').replace('_x000D_', ' ')).strip()


def righe_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    nome = [n for n in wb.sheetnames if n.strip().lower() == 'ore']
    if not nome:
        sys.exit('nel file non c\'e\' un foglio «Ore» (fogli: %s)' % ', '.join(wb.sheetnames))
    it = wb[nome[0]].iter_rows(values_only=True)
    return [str(h or '').strip() for h in next(it)], it


def righe_csv(path):
    testo = open(path, encoding='utf-8-sig').read()
    sep = ';' if testo.split('\n', 1)[0].count(';') > testo.split('\n', 1)[0].count(',') else ','
    rd = csv.reader(io.StringIO(testo), delimiter=sep)
    return [h.strip() for h in next(rd)], rd


def righe_pkl(path):
    D = pickle.load(open(path, 'rb'))
    return [str(h or '').strip() for h in D['hdr']], iter(D['rows'])


def aggrega(hdr, rows):
    manca = [c for c in ('Anno', 'Mese', 'Progetto', 'Reparto', 'Ditta', 'Ore (pivot)') if c not in hdr]
    if manca:
        sys.exit('foglio «Ore» senza le colonne: %s' % ', '.join(manca))
    i = {h: k for k, h in enumerate(hdr)}

    def g(r, c):
        k = i.get(c)
        return r[k] if k is not None and k < len(r) else None
    P = defaultdict(lambda: {'uff': 0.0, 'off': 0.0, 'est': 0.0, 'estEur': 0.0, 'estSU': 0.0, 'estSO': 0.0, 'valG': 0.0, 'mesi': defaultdict(float),
                             'dip': defaultdict(float), 'att': defaultdict(float), 'nd': defaultdict(lambda: [0.0, 0.0, 0.0])})
    tot = {'righe': 0, 'ore': 0.0, 'val': 0.0, 'reparti': set(), 'ditte': set(), 'ultima': None, 'senza_progetto': 0.0}
    for r in rows:
        if not r or all(x in (None, '') for x in r):
            continue
        tot['righe'] += 1
        ore = num(g(r, 'Ore (pivot)'))
        grezzo = g(r, 'Valorizzazione')
        val = num(grezzo)
        senza_val = grezzo in (None, '') or str(grezzo).strip().startswith('#') or not val
        rep = str(g(r, 'Reparto') or '').strip().upper()
        ditta = str(g(r, 'Ditta') or '').strip().upper()
        tot['reparti'].add(rep)
        tot['ditte'].add(ditta)
        tot['ore'] += ore
        tot['val'] += val
        try:
            a, m = int(num(g(r, 'Anno'))), int(num(g(r, 'Mese')))
            d = dt.date(a, m, int(num(g(r, 'Giorno'))) or 1)
        except (ValueError, TypeError):
            try:
                d = dt.date(a, m, 1)
            except (ValueError, TypeError, NameError, UnboundLocalError):
                a = m = 0
                d = None
        p = norm_code(g(r, 'Progetto'))
        if not p:
            tot['senza_progetto'] += ore
            continue
        if not ore and not val:
            continue
        if d and (tot['ultima'] is None or d > tot['ultima']):
            tot['ultima'] = d
        x = P[p]
        esterna = ditta not in ('P&B', '')
        if esterna:
            x['est'] += ore
            x['estEur'] += val
            if senza_val and ore:
                x['estSU' if rep in REP_UFF else 'estSO'] += ore
                tot['est_senza_val'] = tot.get('est_senza_val', 0.0) + ore
        elif rep in REP_UFF:
            x['uff'] += ore
        elif rep in REP_OFF:
            x['off'] += ore
        else:
            x['off'] += ore  # reparto sconosciuto di P&B: produzione, e lo segnalo nel resoconto
            tot.setdefault('rep_ignoti', set()).add(rep)
        x['valG'] += val
        if a and m:
            x['mesi']['%04d-%02d' % (a, m)] += ore
        dip = pulito(g(r, 'Dipendente'))
        if dip:
            x['dip'][dip] += ore
        x['att'][pulito(g(r, 'Attività'))] += ore
        nd = pulito(g(r, 'Nodo')) or '-'
        if nd.upper() == p or nd.upper() in ('MANCANTE', '#N/A'):
            nd = '-'  # nodo uguale al codice della commessa = senza nodo, come nel consuntivo
        k = 2 if esterna else (0 if rep in REP_UFF else 1)
        x['nd'][nd][k] += ore
    return P, tot


def r1(v):
    return round(v, 1) if abs(v - round(v)) > 1e-9 else int(round(v))


def voce_ore(x, src):
    mesi = sorted((k, r1(v)) for k, v in x['mesi'].items() if round(v, 1))
    top = lambda d, n: [[k, r1(v)] for k, v in sorted(d.items(), key=lambda kv: -kv[1])[:n] if round(v, 1)]
    nd = sorted(x['nd'].items(), key=lambda kv: -sum(kv[1]))[:15]
    return {'uff': r1(x['uff']), 'off': r1(x['off']), 'est': r1(x['est']), 'estEur': round(x['estEur']),
            'estSU': r1(x['estSU']), 'estSO': r1(x['estSO']),
            'valG': round(x['valG']), 'first': mesi[0][0] if mesi else None, 'last': mesi[-1][0] if mesi else None,
            'storia': mesi, 'dip': top(x['dip'], 8), 'att': top(x['att'], 8),
            'nd': [[k, r1(v[0]), r1(v[1]), r1(v[2])] for k, v in nd if round(sum(v), 1)], 'src': src}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument('--xlsx')
    g.add_argument('--csv')
    g.add_argument('--pkl')
    ap.add_argument('--state', required=True)
    ap.add_argument('--out')
    ap.add_argument('--file', help='nome del file del gestionale da citare nella pagina')
    ap.add_argument('--dry', action='store_true')
    ap.add_argument('--forza', action='store_true', help='accetta anche un\'estrazione piu\' vecchia dell\'ultimo caricamento')
    ap.add_argument('--min-righe', type=int, default=1000,
                    help='righe minime del foglio (la routine passa il 95%% delle righe dell\'ultimo caricamento buono: un foglio arrivato a meta\' non passa)')
    a = ap.parse_args()
    if a.xlsx:
        hdr, rows = righe_xlsx(a.xlsx)
    elif a.csv:
        hdr, rows = righe_csv(a.csv)
    else:
        hdr, rows = righe_pkl(a.pkl)
    P, tot = aggrega(hdr, rows)
    # un foglio rotto non deve cancellare le ore buone
    rotto = []
    if tot['righe'] < max(1000, a.min_righe):
        rotto.append('solo %d righe (minimo %d)' % (tot['righe'], max(1000, a.min_righe)))
    if tot['ore'] <= 0:
        rotto.append('zero ore in totale')
    if len(tot['reparti'] - {''}) < 2:
        rotto.append('un solo reparto (%s)' % ', '.join(sorted(tot['reparti'])))
    if tot['val'] <= 0:
        rotto.append('valorizzazione tutta vuota o in errore')
    if rotto:
        sys.exit('foglio «Ore» inutilizzabile (%s): lo stato NON viene toccato. Chiedere a Luca di rifare l\'estrazione.' % '; '.join(rotto))
    S = json.load(open(a.state, encoding='utf-8'))
    ultima = tot['ultima'].isoformat() if tot['ultima'] else None
    prima = (S.get('oreG') or {}).get('agg')
    if prima and ultima and ultima < prima and not a.forza:
        sys.exit('estrazione con ore fino al %s, piu\' vecchia dell\'ultimo caricamento (%s): lo stato NON viene toccato '
                 '(--forza per tornare indietro apposta).' % (ultima, prima))
    nome = a.file or (a.xlsx or a.csv or a.pkl).split('/')[-1]
    src = 'Estrazione gestionale %s · ore registrate fino al %s' % (nome, '%s/%s/%s' % (ultima[8:10], ultima[5:7], ultima[:4]) if ultima else '—')
    # alias indicati da Luca: {codice commessa: [codici progetto del foglio Ore da sommare]}, conservati nello stato
    alias = {norm_code(k): [norm_code(x) for x in (v if isinstance(v, list) else [v]) if norm_code(x)]
             for k, v in ((S.get('oreG') or {}).get('alias') or {}).items()}
    codici = {norm_code(c.get('code')) for c in S.get('commesse') or []}
    orig = {q: copy.deepcopy(P[q]) for q in {q for v in alias.values() for q in v} if q in P}  # i valori prima di ogni somma: niente catene
    usati, saltati = set(), []
    for code, altri in alias.items():
        for q in dict.fromkeys(altri):  # un codice ripetuto conta una volta
            if q == code or q not in orig:
                continue
            if q in codici:
                saltati.append([code, q, 'e\' una commessa della pagina: ha gia\' le sue ore'])
                continue
            if q in usati:
                saltati.append([code, q, 'gia\' sommato a un\'altra commessa'])
                continue
            x, y = P[code], orig[q]
            for f in ('uff', 'off', 'est', 'estEur', 'estSU', 'estSO', 'valG'):
                x[f] += y[f]
            for f in ('mesi', 'dip', 'att'):
                for kk, vv in y[f].items():
                    x[f][kk] += vv
            for kk, vv in y['nd'].items():
                for j in range(3):
                    x['nd'][kk][j] += vv[j]
            usati.add(q)
    stat = {'aperte': [0, 0], 'sospese': [0, 0], 'evase': [0, 0]}
    senza = []
    for c in S.get('commesse') or []:
        code = norm_code(c.get('code'))
        chi = 'evase' if c.get('ev') else 'sospese' if c.get('sp') else 'aperte'
        stat[chi][1] += 1
        x = P.get(code)
        if x and (x['uff'] or x['off'] or x['est']):
            c['ore'] = voce_ore(x, src)
            stat[chi][0] += 1
        else:
            c['ore'] = {'uff': None, 'off': None, 'est': None, 'storia': [], 'src': 'Nessuna ora registrata su questa commessa: ' + src}
            senza.append((chi, c.get('code')))
    fuori = sorted(((p, r1(x['uff'] + x['off'] + x['est'])) for p, x in P.items() if p not in codici and p not in usati), key=lambda t: -t[1])
    S['oreG'] = {'agg': ultima, 'alias': (S.get('oreG') or {}).get('alias') or {}, 'file': nome + ' · foglio «Ore»',
                 'fonte': 'ore per commessa: uff = reparti UFF+TEC ditta P&B; off = reparti OFF+ART ditta P&B; est/estEur = righe con ditta diversa da P&B; valG = valorizzazione totale. Aggiornate su tutte le commesse (aperte, sospese, evase).',
                 'controllo': {'righe': tot['righe'], 'ore': round(tot['ore']), 'valorizzazione': round(tot['val']),
                               'reparti': sorted(tot['reparti']), 'ditte': sorted(tot['ditte']), 'progetti': len(P),
                               'oreSenzaProgetto': round(tot['senza_progetto']),
                               'oreEsterneSenzaValorizzazione': round(tot.get('est_senza_val', 0)),
                               'aliasSaltati': saltati,
                               'progettiNonInPagina': [[p, h] for p, h in fuori[:30]],
                               'repartiSconosciuti': sorted(tot.get('rep_ignoti', set()))}}
    print('foglio «Ore»: %d righe, %d ore, valorizzazione %d €, %d progetti, ultima registrazione %s'
          % (tot['righe'], tot['ore'], tot['val'], len(P), ultima))
    for k in ('aperte', 'sospese', 'evase'):
        print('  commesse %-7s con ore %d su %d' % (k, stat[k][0], stat[k][1]))
    print('  senza ore: %d (%s%s)' % (len(senza), ', '.join(c for _, c in senza[:12]), '…' if len(senza) > 12 else ''))
    print('  progetti del gestionale non presenti nella pagina: %d (%s)' % (len(fuori), ', '.join('%s %sh' % t for t in fuori[:8])))
    if not a.dry:
        json.dump(S, open(a.out or a.state, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))


if __name__ == '__main__':
    main()
