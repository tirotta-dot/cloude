#!/usr/bin/env python3
"""Genera OFFER_FS-11-17_Control_System_EN.docx dal file markdown omonimo."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

HERE = Path(__file__).parent
SRC = HERE / "OFFER_FS-11-17_Control_System_EN.md"
DST = HERE / "OFFER_FS-11-17_Control_System_EN.docx"

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def add_inline(paragraph, text):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Cm(2))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            h = doc.add_heading("", level=min(level, 3))
            add_inline(h, BOLD_RE.sub(r"\1", line.lstrip("# ").strip()))
            if level == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    p = table.rows[r].cells[c].paragraphs[0]
                    add_inline(p, cell)
                    if r == 0:
                        for run in p.runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        if line.lstrip().startswith("- "):
            indent = (len(line) - len(line.lstrip())) // 2
            style = "List Bullet" if indent == 0 else "List Bullet 2"
            p = doc.add_paragraph(style=style)
            add_inline(p, line.lstrip()[2:].strip())
            i += 1
            continue
        # paragrafo: accorpa le righe consecutive dello stesso blocco
        block = [line]
        while (i + 1 < len(lines) and lines[i + 1].strip()
               and not lines[i + 1].startswith(("#", "|", "---"))
               and not lines[i + 1].lstrip().startswith("- ")):
            i += 1
            block.append(lines[i].rstrip())
        p = doc.add_paragraph()
        for n, part in enumerate(block):
            if n:
                p.add_run().add_break()
            add_inline(p, part)
        i += 1

    doc.save(DST)
    print(f"scritto {DST}")


if __name__ == "__main__":
    main()
